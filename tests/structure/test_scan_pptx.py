"""PLAN-030g：PPTX 双模式扫描与嵌图执行。"""
from __future__ import annotations

import hashlib
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation

from qyunslation.exporter.pptx.pptx2html_exporter import PPTX2HTMLExporterConfig
from qyunslation.ir.document import Document
from qyunslation.structure import ManifestStore, PptxStructureScanner
from qyunslation.structure.ingest import InputPreparationError
from qyunslation.structure.models import (
    CURRENT_SCHEMA_VERSION,
    CanvasKind,
    ExecutionStatus,
    ObjectType,
    OutputEditability,
    ProcessingMode,
)
from qyunslation.structure.scan_pptx import pack_image_pptx
from qyunslation.translator.ai_translator.pptx_translator import (
    PPTXTranslator,
    PPTXTranslatorConfig,
)
from qyunslation.workflow.pptx_workflow import PPTXWorkflow, PPTXWorkflowConfig


def _png_stub() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc``\x00\x00"
        b"\x00\x04\x00\x01\xf6\x178U\x00\x00\x00\x00IEND\xaeB`\x82"
    )


def test_scan_pptx_emits_slide_objects(generated_structure_fixtures: Path):
    manifest = PptxStructureScanner().scan(generated_structure_fixtures / "presentation.pptx")

    assert manifest.schema_version == CURRENT_SCHEMA_VERSION
    assert len(manifest.canvases) == 1
    assert manifest.canvases[0].kind is CanvasKind.SLIDE
    images = [obj for obj in manifest.objects if obj.type is ObjectType.IMAGE]
    tables = [obj for obj in manifest.objects if obj.type is ObjectType.TABLE]
    textboxes = [obj for obj in manifest.objects if obj.type is ObjectType.TEXT_BOX]
    assert len(images) == 1
    assert images[0].execution_status is ExecutionStatus.PENDING
    assert images[0].source_refs[0].kind.value == "PPTX_SHAPE"
    assert len(tables) == 1
    assert textboxes
    assert manifest.document.output_editability is OutputEditability.EDITABLE


def test_native_roundtrip_keeps_slide_geometry(generated_structure_fixtures: Path):
    source = generated_structure_fixtures / "presentation.pptx"
    before = Presentation(BytesIO(source.read_bytes()))
    document = Document.from_bytes(source.read_bytes(), suffix=".pptx", stem=source.stem)
    translator = PPTXTranslator(PPTXTranslatorConfig(skip_translate=True))
    translator.translate(document)
    after = Presentation(BytesIO(document.content))

    assert len(after.slides) == len(before.slides)
    assert after.slide_width == before.slide_width
    assert after.slide_height == before.slide_height
    assert after.slide_width > 0


def test_rendered_scan_uses_injected_renderer(generated_structure_fixtures: Path):
    source = generated_structure_fixtures / "presentation.pptx"
    page = _png_stub()
    scanner = PptxStructureScanner(slide_renderer=lambda _content: [page])
    manifest = scanner.scan(source, processing_mode=ProcessingMode.RENDERED)

    images = [obj for obj in manifest.objects if obj.type is ObjectType.IMAGE]
    assert len(images) == 1
    assert images[0].canvas_id == "slide:1"
    assert manifest.document.selected_mode is ProcessingMode.RENDERED
    assert manifest.document.output_editability is OutputEditability.RASTERIZED
    assert manifest.canvases[0].reading_order == [images[0].object_id]


def test_rendered_missing_renderer_fails_closed(monkeypatch, generated_structure_fixtures: Path):
    monkeypatch.setattr("qyunslation.structure.scan_pptx.shutil.which", lambda _name: None)
    with pytest.raises(InputPreparationError) as exc_info:
        PptxStructureScanner().scan(
            generated_structure_fixtures / "presentation.pptx",
            processing_mode=ProcessingMode.RENDERED,
        )
    assert exc_info.value.code == "RUNTIME_CAPABILITY_UNAVAILABLE"


def test_pack_image_pptx_preserves_page_order():
    pages = [_png_stub(), _png_stub()]
    packed = pack_image_pptx(pages, width_pt=720.0, height_pt=405.0)
    presentation = Presentation(BytesIO(packed))
    assert len(presentation.slides) == 2
    assert presentation.slide_width > 0


def test_workflow_writes_image_execution(
    generated_structure_fixtures: Path, tmp_path, monkeypatch
):
    monkeypatch.setenv("QYUNSLATION_MANIFEST_CACHE", str(tmp_path))
    source = generated_structure_fixtures / "presentation.pptx"
    workflow = PPTXWorkflow(
        config=PPTXWorkflowConfig(
            translator_config=PPTXTranslatorConfig(skip_translate=True),
            html_exporter_config=PPTX2HTMLExporterConfig(),
        )
    )
    workflow.read_bytes(source.read_bytes(), stem=source.stem, suffix=".pptx")
    monkeypatch.setattr(
        "qyunslation.extensions.image_translate.translate_image_bytes",
        lambda data, suffix, to_lang: (data, 2, {}),
    )
    workflow.translate()
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    updated = ManifestStore().get_execution(digest)
    assert updated is not None
    images = [obj for obj in updated.objects if obj.type is ObjectType.IMAGE]
    assert images
    assert images[0].execution_status in {
        ExecutionStatus.TRANSLATED,
        ExecutionStatus.EXPLICITLY_SKIPPED,
        ExecutionStatus.FAILED_SOFT,
    }
    after = Presentation(BytesIO(workflow.document_translated.content))
    assert len(after.slides) == 1


def test_media_hash_survives_skip_translate(generated_structure_fixtures: Path):
    source = generated_structure_fixtures / "presentation.pptx"
    before = _media_hashes(source.read_bytes())
    document = Document.from_bytes(source.read_bytes(), suffix=".pptx", stem=source.stem)
    PPTXTranslator(PPTXTranslatorConfig(skip_translate=True)).translate(document)
    assert _media_hashes(document.content) == before


def _media_hashes(content: bytes) -> dict[str, str]:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        names = sorted(
            name
            for name in archive.namelist()
            if name.startswith("ppt/media/") and not name.endswith("/")
        )
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in names}
