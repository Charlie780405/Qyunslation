"""PLAN-037 P1-B：PPTX 表格 policy 执行分流。"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from qyunslation.structure import PptxStructureScanner
from qyunslation.structure.models import TranslationPolicy
from qyunslation.structure.pptx_table_exec import merge_docx_translations, partition_pptx_segments
from qyunslation.translator.ai_translator.pptx_translator import PPTXTranslator, PPTXTranslatorConfig


def _table_pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1.5)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Rate"
    table.cell(1, 1).text = "42.3%"
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def test_partition_skips_numeric_pptx_table_cell():
    content = _table_pptx_bytes()
    manifest = PptxStructureScanner().scan(content, source_name="table-policy.pptx")
    translator = PPTXTranslator(PPTXTranslatorConfig(skip_translate=True))
    prs, elements, originals = translator._pre_translate(
        type("Doc", (), {"content": content})()
    )
    text_elements = [info for info in elements if info.get("type") != "image"]
    batch = partition_pptx_segments(originals, text_elements, manifest)
    assert "42.3%" not in batch.llm_texts
    merged = merge_docx_translations(
        originals,
        [f"译:{text}" for text in batch.llm_texts],
        batch.llm_to_original,
        batch.token_maps,
        batch.preserved_indices,
    )
    assert merged[originals.index("42.3%")] == "42.3%"
