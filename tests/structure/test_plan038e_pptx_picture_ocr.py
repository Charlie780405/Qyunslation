# SPDX-License-Identifier: MPL-2.0
"""PLAN-038e：PPTX picture OCR 嵌字链证明。"""
from __future__ import annotations

import hashlib
import zipfile
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from qyunslation.structure.models import ExecutionStatus, ObjectType
from qyunslation.translator.ai_translator.pptx_translator import (
    PPTXTranslator,
    PPTXTranslatorConfig,
)


def _png() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc``\x00\x00"
        b"\x00\x04\x00\x01\xf6\x178U\x00\x00\x00\x00IEND\xaeB`\x82"
    )


def _pptx_with_picture(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(_png()), Inches(1), Inches(1), width=Inches(1))
    out = tmp_path / "pic.pptx"
    prs.save(out)
    return out


def _media_hashes(content: bytes) -> dict[str, str]:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        names = sorted(
            name
            for name in archive.namelist()
            if name.startswith("ppt/media/") and not name.endswith("/")
        )
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in names}


def test_pptx_picture_overlay_replaces_blob_and_marks_translated(tmp_path: Path, monkeypatch):
    source = _pptx_with_picture(tmp_path)
    before = source.read_bytes()
    before_hashes = _media_hashes(before)
    replaced = bytearray(_png())
    replaced[-8] = (replaced[-8] + 1) % 256
    new_blob = bytes(replaced)

    monkeypatch.setattr(
        "qyunslation.extensions.image_translate.translate_image_bytes",
        lambda data, suffix, to_lang: (new_blob, 3, {"ok": True}),
    )

    translator = PPTXTranslator(PPTXTranslatorConfig(skip_translate=True))
    translator.skip_translate = False  # exercise overlay path without LLM
    prs = Presentation(BytesIO(before))
    elements = []
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                elements.append(
                    {
                        "type": "image",
                        "shape": shape,
                        "slide_index": index,
                        "shape_id": shape.shape_id,
                        "execution_status": "PENDING",
                    }
                )
    assert elements
    translator._overlay_images(elements)
    assert elements[0]["execution_status"] == "TRANSLATED"

    out = BytesIO()
    prs.save(out)
    after_hashes = _media_hashes(out.getvalue())
    assert after_hashes != before_hashes

    # Manifest writeback path
    from qyunslation.structure.scan_pptx import PptxStructureScanner

    manifest = PptxStructureScanner().scan(before, source_name="pic.pptx")
    translator._writeback_manifest(manifest, elements)
    images = [o for o in manifest.objects if o.type is ObjectType.IMAGE]
    assert images
    assert images[0].execution_status is ExecutionStatus.TRANSLATED
