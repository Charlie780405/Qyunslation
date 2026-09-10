# SPDX-License-Identifier: MPL-2.0
"""PLAN-038g：DOCX 跨页/续表 caption → 同 semantic_id 多 occurrence。"""
from __future__ import annotations

from io import BytesIO

import docx

from qyunslation.structure import DocxStructureScanner
from qyunslation.structure.models import ObjectType


def _docx_continued_tables() -> bytes:
    document = docx.Document()
    document.add_paragraph("Table 1. Baseline characteristics")
    t1 = document.add_table(rows=2, cols=2)
    t1.cell(0, 0).text = "Arm"
    t1.cell(0, 1).text = "N"
    t1.cell(1, 0).text = "A"
    t1.cell(1, 1).text = "10"
    document.add_paragraph("Table 1 Continued")
    t2 = document.add_table(rows=2, cols=2)
    t2.cell(0, 0).text = "Arm"
    t2.cell(0, 1).text = "N"
    t2.cell(1, 0).text = "B"
    t2.cell(1, 1).text = "12"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_docx_continued_table_shares_semantic_id_with_occurrences():
    manifest = DocxStructureScanner().scan(
        _docx_continued_tables(), source_name="continued.docx"
    )
    tables = [o for o in manifest.objects if o.type is ObjectType.TABLE]
    assert len(tables) == 2
    assert tables[0].semantic_id == "table:1"
    assert tables[1].semantic_id == "table:1"
    assert tables[0].semantic_occurrence_index == 1
    assert tables[1].semantic_occurrence_index == 2
    # row indexes continue across occurrences
    rows0 = [b.row_index for b in tables[0].translatable_blocks if b.row_index is not None]
    rows1 = [b.row_index for b in tables[1].translatable_blocks if b.row_index is not None]
    assert max(rows0) < min(rows1)
