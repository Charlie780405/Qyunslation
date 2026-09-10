"""PLAN-036：DOCX 表格 policy 执行分流。"""
from __future__ import annotations

from io import BytesIO

import docx

from qyunslation.structure import DocxStructureScanner
from qyunslation.structure.docx_table_exec import merge_docx_translations, partition_docx_segments
from qyunslation.structure.docx_walk import walk_docx
from qyunslation.structure.models import TranslationPolicy


def _numeric_table_docx() -> bytes:
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Rate"
    table.cell(1, 1).text = "42.3%"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_partition_skips_preserve_cells_and_protects_mixed_cells():
    content = _numeric_table_docx()
    manifest = DocxStructureScanner().scan(content, source_name="numeric.docx")
    _doc, _segments, elements, originals = walk_docx(content)
    batch = partition_docx_segments(originals, elements, manifest)

    assert len(batch.llm_texts) == 3
    assert "42.3%" not in batch.llm_texts
    assert batch.preserved_indices

    fake_llm = [f"译:{text}" for text in batch.llm_texts]
    merged = merge_docx_translations(
        originals,
        fake_llm,
        batch.llm_to_original,
        batch.token_maps,
        batch.preserved_indices,
    )
    assert merged[originals.index("42.3%")] == "42.3%"
    assert merged[originals.index("Rate")].startswith("译:")


def test_fallback_classifies_cells_without_manifest():
    content = _numeric_table_docx()
    _doc, _segments, elements, originals = walk_docx(content)
    batch = partition_docx_segments(originals, elements, None)
    preserved_texts = {originals[i] for i in batch.preserved_indices}
    assert "42.3%" in preserved_texts


def test_manifest_policy_map_matches_table_cells():
    content = _numeric_table_docx()
    manifest = DocxStructureScanner().scan(content, source_name="numeric.docx")
    table = next(obj for obj in manifest.objects if obj.type.value == "TABLE")
    by_text = {block.source_text: block.translation_policy for block in table.translatable_blocks}
    assert by_text["42.3%"] is TranslationPolicy.PRESERVE
    assert by_text["Rate"] is TranslationPolicy.TRANSLATE
