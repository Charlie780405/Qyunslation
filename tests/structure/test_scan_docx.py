"""PLAN-030e：DOCX 结构扫描与结构校验门。"""
from __future__ import annotations

import hashlib
import zipfile
from io import BytesIO
from pathlib import Path

import docx
import pytest

from qyunslation.extensions.docx_image_overlay import enumerate_drawing_occurrences
from qyunslation.structure import DocxStructureScanner, ManifestStore
from qyunslation.structure.docx_walk import walk_docx
from qyunslation.structure.models import (
    CURRENT_SCHEMA_VERSION,
    CanvasKind,
    ExecutionStatus,
    ObjectType,
)

ROOT = Path(__file__).resolve().parents[2]
REVIEW = ROOT / "tests/fixtures/structure/review.docx"


@pytest.fixture
def review_docx(generated_structure_fixtures: Path) -> Path:
    path = generated_structure_fixtures / "review.docx"
    assert path.is_file()
    return path


def test_scan_docx_produces_flow_layout_objects_without_bbox(review_docx: Path):
    manifest = DocxStructureScanner().scan(review_docx)

    assert manifest.schema_version == CURRENT_SCHEMA_VERSION
    assert manifest.canvases[0].kind is CanvasKind.SECTION
    bodies = [obj for obj in manifest.objects if obj.type is ObjectType.BODY]
    assert bodies
    assert all(obj.bbox is None for obj in bodies)


def test_table_has_cell_level_blocks_and_dimensions(review_docx: Path):
    manifest = DocxStructureScanner().scan(review_docx)
    tables = [obj for obj in manifest.objects if obj.type is ObjectType.TABLE]

    assert len(tables) == 1
    table = tables[0]
    assert table.row_count == 2
    assert table.column_count == 2
    assert len(table.translatable_blocks) == 4


def test_textbox_header_footer_and_multi_section_fixture(review_docx: Path):
    manifest = DocxStructureScanner().scan(review_docx)
    counts = manifest.summary.object_counts

    assert counts.get("TEXT_BOX", 0) >= 1
    assert counts.get("CAPTION", 0) >= 2
    assert counts.get("FIGURE", 0) + counts.get("IMAGE", 0) == 2

    _, segments, _, _ = walk_docx(review_docx.read_bytes())
    containers = {segment.container_ref for segment in segments}
    assert any("header" in ref for ref in containers)
    assert any("footer" in ref for ref in containers)
    assert any(ref.startswith("section:2/body") for ref in containers)


def test_second_section_body_maps_to_section_two_canvas(review_docx: Path):
    manifest = DocxStructureScanner().scan(review_docx)
    second_bodies = [
        obj
        for obj in manifest.objects
        if obj.type is ObjectType.BODY
        and any(
            "Second section body" in block.source_text
            for block in obj.translatable_blocks
        )
    ]
    assert len(second_bodies) == 1
    assert second_bodies[0].canvas_id == "section:2"
    assert second_bodies[0].source_refs[0].ref.startswith("section:2/")


def test_structural_integrity_after_roundtrip_read(review_docx: Path):
    before = _structure_fingerprint(review_docx)
    manifest = DocxStructureScanner().scan(review_docx)
    assert manifest.summary.object_counts.get("TABLE") == 1

    doc = docx.Document(str(review_docx))
    assert len(doc.sections) >= 2
    assert len(list(enumerate_drawing_occurrences(doc))) == 2
    after = _structure_fingerprint(review_docx)
    assert before == after


def test_manifest_store_roundtrip_docx(review_docx: Path, tmp_path, monkeypatch):
    monkeypatch.setenv("QYUNSLATION_MANIFEST_CACHE", str(tmp_path))
    manifest = DocxStructureScanner().scan(review_docx)
    store = ManifestStore()
    assert store.put(manifest) is not None
    loaded = store.get(manifest.document.source_sha256)
    assert loaded is not None
    assert loaded.summary.object_counts == manifest.summary.object_counts


def _structure_fingerprint(path: Path) -> dict:
    content = path.read_bytes()
    doc = docx.Document(BytesIO(content))
    tables = doc.tables
    rows_cols = [(len(table.rows), len(table.columns)) for table in tables]
    with zipfile.ZipFile(BytesIO(content)) as archive:
        media = sorted(
            name
            for name in archive.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        )
        media_hashes = {
            name: hashlib.sha256(archive.read(name)).hexdigest() for name in media
        }
    return {
        "sections": len(doc.sections),
        "tables": rows_cols,
        "drawings": len(list(enumerate_drawing_occurrences(doc))),
        "media_hashes": media_hashes,
    }
