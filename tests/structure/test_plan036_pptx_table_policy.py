"""PLAN-036：PPTX 表格 translation_policy。"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from qyunslation.structure import PptxStructureScanner
from qyunslation.structure.models import ObjectType, TranslationPolicy


def test_pptx_table_blocks_carry_translation_policy(tmp_path: Path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1.5)).table
    table.cell(0, 0).text = "Outcome"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Rate"
    table.cell(1, 1).text = "42.3%"
    path = tmp_path / "table-policy.pptx"
    presentation.save(path)

    manifest = PptxStructureScanner().scan(path)
    table_obj = next(obj for obj in manifest.objects if obj.type is ObjectType.TABLE)
    by_text = {block.source_text: block for block in table_obj.translatable_blocks}
    assert by_text["42.3%"].translation_policy is TranslationPolicy.PRESERVE
    assert by_text["Rate"].translation_policy is TranslationPolicy.TRANSLATE
    assert by_text["Rate"].role is not None
