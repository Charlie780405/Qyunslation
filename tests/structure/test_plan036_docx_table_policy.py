"""PLAN-036b：DOCX 表格 translation_policy。"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import docx
import pytest

from qyunslation.structure import DocxStructureScanner
from qyunslation.structure.models import ObjectType, TranslationPolicy

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
def review_docx(generated_structure_fixtures: Path) -> Path:
    path = generated_structure_fixtures / "review.docx"
    assert path.is_file()
    return path


def test_docx_table_blocks_carry_translation_policy(review_docx: Path):
    manifest = DocxStructureScanner().scan(review_docx)
    table = next(obj for obj in manifest.objects if obj.type is ObjectType.TABLE)
    assert table.translatable_blocks
    for block in table.translatable_blocks:
        assert block.translation_policy is not None
        assert block.role is not None


def test_docx_numeric_cell_is_preserve():
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Rate"
    table.cell(1, 1).text = "42.3%"
    stream = BytesIO()
    document.save(stream)

    manifest = DocxStructureScanner().scan(stream.getvalue(), source_name="numeric.docx")
    table_obj = next(obj for obj in manifest.objects if obj.type is ObjectType.TABLE)
    by_text = {b.source_text: b for b in table_obj.translatable_blocks}
    assert by_text["42.3%"].translation_policy is TranslationPolicy.PRESERVE
    assert by_text["Rate"].translation_policy is TranslationPolicy.TRANSLATE
