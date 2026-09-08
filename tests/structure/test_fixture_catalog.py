from __future__ import annotations

import hashlib
import importlib.util
import json
import re
from pathlib import Path

from docx import Document
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from qyunslation.structure import ContentProfile, SourceFormat


ROOT = Path(__file__).resolve().parents[2]
FIXTURE_ROOT = ROOT / "tests" / "fixtures" / "structure"
CATALOG_PATH = FIXTURE_ROOT / "catalog.v1.json"
TRUTH_PATH = FIXTURE_ROOT / "ljae439.truth.json"
GENERATOR_PATH = FIXTURE_ROOT / "generate_synthetic.py"
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")

CORE_FORMATS = {
    SourceFormat.PDF,
    SourceFormat.DOCX,
    SourceFormat.PNG,
    SourceFormat.JPEG,
    SourceFormat.WEBP,
    SourceFormat.BMP,
    SourceFormat.TIFF,
    SourceFormat.PPTX,
}
REQUIRED_PROFILES = {
    ContentProfile.RESEARCH_ARTICLE,
    ContentProfile.REVIEW_ARTICLE,
    ContentProfile.PRESENTATION,
    ContentProfile.POSTER,
}


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _load_generator():
    spec = importlib.util.spec_from_file_location("plan030_fixtures", GENERATOR_PATH)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_catalog_covers_core_formats_and_priority_content_profiles():
    catalog = _read_json(CATALOG_PATH)
    fixtures = catalog["fixtures"]

    assert catalog["schema_version"] == "1.0.0"
    assert {SourceFormat(item["source_format"]) for item in fixtures} >= CORE_FORMATS
    assert {ContentProfile(item["content_profile"]) for item in fixtures} >= (
        REQUIRED_PROFILES
    )
    assert len({item["id"] for item in fixtures}) == len(fixtures)
    assert all(
        item["future_owner"].startswith("PLAN-030")
        or item["future_owner"].startswith("PLAN-033")
        for item in fixtures
    )


def test_synthetic_fixtures_are_reproducible_and_match_catalog(tmp_path: Path):
    catalog = _read_json(CATALOG_PATH)
    generator = _load_generator()
    first = generator.generate_all(tmp_path / "first")
    second = generator.generate_all(tmp_path / "second")
    expected = {
        item["relative_path"]: item["sha256"]
        for item in catalog["fixtures"]
        if item["origin"] == "GENERATED"
    }

    assert first == second
    assert first == expected
    assert all(SHA256_RE.fullmatch(value) for value in first.values())

    for relative_path, expected_hash in first.items():
        data = (tmp_path / "first" / relative_path).read_bytes()
        assert hashlib.sha256(data).hexdigest() == expected_hash


def test_synthetic_outputs_are_real_parseable_documents(tmp_path: Path):
    generator = _load_generator()
    generator.generate_all(tmp_path)

    for name in ("single-column.pdf", "double-column.pdf"):
        reader = PdfReader(tmp_path / name)
        assert len(reader.pages) == 1
        assert "Figure 1" in reader.pages[0].extract_text()

    docx = Document(tmp_path / "review.docx")
    assert len(docx.tables) == 1
    assert len(docx.inline_shapes) == 2

    presentation = Presentation(tmp_path / "presentation.pptx")
    assert len(presentation.slides) == 1
    assert len(presentation.slides[0].shapes) >= 3

    expected_images = {
        "poster.png": ("PNG", 1),
        "photo.jpg": ("JPEG", 1),
        "diagram.webp": ("WEBP", 1),
        "scan.bmp": ("BMP", 1),
        "multipage.tiff": ("TIFF", 2),
    }
    for name, (format_name, frame_count) in expected_images.items():
        with Image.open(tmp_path / name) as image:
            assert image.format == format_name
            assert getattr(image, "n_frames", 1) == frame_count

    with Image.open(tmp_path / "photo.jpg") as image:
        assert image.getexif()[274] == 6


def test_ljae439_truth_reserves_exact_semantic_ids_for_knowledge_base_import():
    truth = _read_json(TRUTH_PATH)

    assert truth["document_id"] == "doi:10.1093/bjd/ljae439"
    assert truth["scope"] == "MAIN_ARTICLE"
    assert truth["usage_scope"] == "TEST_FIXTURE_ONLY"
    assert truth["source"]["kind"] == "KNOWLEDGE_BASE"
    assert truth["source"]["runtime_dependency"] is False
    assert truth["source"]["locator"].startswith("knowledge-base://")
    assert truth["source"]["import_state"] in {"PENDING", "IMPORTED"}
    assert truth["semantic_ids"]["figures"] == [
        "figure:1",
        "figure:2",
        "figure:3",
        "figure:4",
        "figure:5",
    ]
    assert truth["semantic_ids"]["tables"] == [
        "table:1",
        "table:2",
        "table:3",
    ]

    sha256 = truth["source"]["sha256"]
    if truth["source"]["import_state"] == "IMPORTED":
        assert isinstance(sha256, str) and SHA256_RE.fullmatch(sha256)
    else:
        assert sha256 is None


def test_every_catalog_entry_has_truth_or_deterministic_generation_evidence():
    catalog = _read_json(CATALOG_PATH)

    for item in catalog["fixtures"]:
        if item["origin"] == "GENERATED":
            assert SHA256_RE.fullmatch(item["sha256"])
            assert item["relative_path"]
        elif item["origin"] in {"KNOWLEDGE_BASE", "OPEN_ACCESS"}:
            if item["origin"] == "KNOWLEDGE_BASE":
                assert item["truth_file"] == TRUTH_PATH.relative_to(ROOT).as_posix()
            else:
                assert item["truth_file"]
                assert (ROOT / item["truth_file"]).is_file()
            assert item["fixture_state"] == "MATERIALIZED"
            fixture_path = FIXTURE_ROOT / item["relative_path"]
            assert fixture_path.is_file()
            actual_hash = hashlib.sha256(fixture_path.read_bytes()).hexdigest()
            assert actual_hash == item["sha256"]
        else:
            raise AssertionError(f"unsupported fixture origin: {item['origin']}")
