# SPDX-License-Identifier: MPL-2.0
"""PLAN-060：译后证据不得把参考文献伪装成术语候选。"""
from __future__ import annotations

from PIL import Image

from qyunslation.core.factory import create_workflow_from_payload
from qyunslation.core.schemas import TranslatePayload
from qyunslation.extensions import image_translate
from qyunslation.ir.document import Document
from qyunslation.workflow.image_overlay_workflow import (
    ImageOverlayWorkflow,
    ImageOverlayWorkflowConfig,
)
from qyunslation.workbench import gui_client
from qyunslation.workbench.gui_client import _read_text, build_bilingual_evidence
from qyunslation.workbench.evidence import BilingualTermEvidence, classify_risk, extract_term_pairs


def test_explicit_table_and_ocr_term_evidence_keeps_locator():
    rows = extract_term_pairs(
        [
            BilingualTermEvidence(
                source_text="Primary endpoint",
                target_text="主要终点",
                source_term="primary endpoint",
                target_term="主要终点",
                role="table",
                page_no=3,
                block_id="table-1-r2-c1",
                object_id="table-1",
            ),
            BilingualTermEvidence(
                source_text="ABC-101",
                target_text="ABC-101",
                source_term="ABC-101",
                target_term="ABC-101",
                role="ocr",
                page_no=4,
                object_id="figure-2",
            ),
        ]
    )

    assert [row["source_term"] for row in rows] == ["primary endpoint", "ABC-101"]
    assert rows[0]["occurrences"][0]["block_id"] == "table-1-r2-c1"
    assert rows[1]["occurrences"][0]["object_id"] == "figure-2"


def test_reference_evidence_is_never_extracted_as_a_candidate():
    rows = extract_term_pairs(
        [
            BilingualTermEvidence(
                source_text="ABC-101", target_text="ABC-101", role="reference"
            )
        ]
    )
    assert rows == []


def test_unaligned_code_candidate_does_not_fabricate_a_target_translation():
    rows = extract_term_pairs(
        [BilingualTermEvidence(source_text="ABC-101 was administered.", target_text="已给药。")]
    )
    assert rows == [
        {
            "source_term": "ABC-101",
            "observed_target": "",
            "term_type": "code",
            "source_context": "ABC-101 was administered.",
            "target_context": "已给药。",
            "occurrences": [
                {
                    "page_no": None,
                    "block_id": None,
                    "object_id": None,
                    "char_start": None,
                    "char_end": None,
                    "bbox": None,
                    "source_context": "ABC-101 was administered.",
                    "target_context": "已给药。",
                }
            ],
        }
    ]


def test_image_ocr_evidence_requires_geometric_alignment(monkeypatch, tmp_path):
    source, target = tmp_path / "source.png", tmp_path / "target.png"
    Image.new("RGB", (200, 100), "white").save(source)
    Image.new("RGB", (200, 100), "white").save(target)

    def fake_ocr(path):
        if str(path) == str(source):
            return [(10, 10, 90, 30, "Primary endpoint", 0.99)]
        return [(12, 10, 90, 30, "主要终点", 0.99)]

    monkeypatch.setattr("qyunslation.extensions.image_translate.ocr_image", fake_ocr)
    rows, reason = build_bilingual_evidence(source, target)

    assert reason is None
    assert len(rows) == 1
    assert rows[0].role == "ocr"
    assert rows[0].source_term == "Primary endpoint"
    assert rows[0].target_term == "主要终点"


def test_high_risk_heuristics_cover_embedded_drug_target_protocol_and_organization_terms():
    assert classify_risk("dupilumab treatment") == "high"
    assert classify_risk("IL-4Rα antibody") == "high"
    assert classify_risk("SSGJ-611-CRS-III-01") == "high"
    assert classify_risk("Capital Medical University") == "high"


def test_office_workflows_keep_hard_terms_from_the_term_policy_payload():
    from pydantic import TypeAdapter

    glossary = {"primary endpoint": "主要终点评估"}
    for workflow_type in ("docx", "pptx"):
        payload = TypeAdapter(TranslatePayload).validate_python(
            {
                "workflow_type": workflow_type,
                "skip_translate": True,
                "glossary_dict": glossary,
                "termbase_policy": {"schema": "058-term-policy-v1", "terms": []},
                "termbase_version": "058-test",
            }
        )
        workflow = create_workflow_from_payload(payload)
        assert workflow.config.translator_config.glossary_dict == glossary


def test_image_ocr_receives_hard_terms_from_the_current_translation_policy(monkeypatch):
    seen = {}

    def fake_translate(data, *, suffix, to_lang, glossary=None):
        seen.update({"suffix": suffix, "to_lang": to_lang, "glossary": glossary})
        return data, 0, {}

    monkeypatch.setattr(
        "qyunslation.extensions.image_translate.translate_image_bytes", fake_translate
    )
    workflow = ImageOverlayWorkflow(
        config=ImageOverlayWorkflowConfig(
            to_lang="简体中文",
            glossary_dict={"primary endpoint": "主要终点评估"},
        )
    )
    workflow.read_bytes(b"not-an-image", stem="figure", suffix=".png")
    workflow.translate()

    assert seen["suffix"] == ".png"
    assert seen["to_lang"] == "简体中文"
    assert seen["glossary"] == {"primary endpoint": "主要终点评估"}


def test_task_term_policy_overrides_static_glossary_for_image_ocr(monkeypatch):
    prompts = []

    monkeypatch.setattr(
        image_translate,
        "_load_glossary",
        lambda: {"primary endpoint": "过时译法"},
    )

    def fake_chat(prompt, _model, _num_ctx):
        prompts.append(prompt)
        return "1. 主要终点评估"

    monkeypatch.setattr(image_translate, "_chat", fake_chat)

    translated = image_translate.translate_texts(
        ["primary endpoint"],
        glossary={"primary endpoint": "主要终点评估"},
    )

    assert translated == {1: "主要终点评估"}
    assert "1. 主要终点评估" in prompts[0]


def test_image_source_uses_ocr_text_before_term_policy_is_resolved(monkeypatch, tmp_path):
    source = tmp_path / "figure.png"
    source.write_bytes(b"image")
    monkeypatch.setattr(
        "qyunslation.extensions.image_translate.ocr_image",
        lambda _path: [
            (0, 0, 100, 20, "primary endpoint", 0.99),
            (0, 30, 100, 50, "dupilumab", 0.99),
        ],
    )

    assert _read_text(source) == "primary endpoint\ndupilumab"


def test_office_source_includes_embedded_image_ocr_before_term_policy(monkeypatch, tmp_path):
    from docx import Document as DocxDocument
    from pptx import Presentation

    docx_path = tmp_path / "source.docx"
    document = DocxDocument()
    document.add_paragraph("正文")
    document.save(docx_path)

    pptx_path = tmp_path / "source.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(pptx_path)

    calls = []

    def fake_package_media(path, prefix):
        calls.append((path.suffix, prefix))
        return ["primary endpoint"]

    monkeypatch.setattr(gui_client, "_package_media_ocr_text", fake_package_media)

    assert "primary endpoint" in _read_text(docx_path)
    assert "primary endpoint" in _read_text(pptx_path)
    assert calls == [(".docx", "word/media/"), (".pptx", "ppt/media/")]
