"""PLAN-078f f2: the pre-check tells people what is limiting about their file, from the file itself."""
from __future__ import annotations

import io
import zipfile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from qyunslation.api.v1 import router as api_v1_router
from qyunslation.persist.db import init_engine, reset_engine
from qyunslation.persist.models import Base
from qyunslation.pipeline import preflight_probe as probe_module
from qyunslation.pipeline.preflight_probe import probe_document

OWNER = {"X-Dev-User": "owner1", "X-Dev-Tenant": "pilot", "X-Dev-Role": "translator"}


@pytest.fixture(autouse=True)
def _probe_on(monkeypatch):
    monkeypatch.setenv("QYUNSLATION_PREFLIGHT_PROBE", "1")


def test_the_probe_can_be_switched_off_by_the_operator(monkeypatch, tmp_path):
    broken = tmp_path / "x.pdf"
    broken.write_bytes(b"this is not a pdf")
    assert probe_document(broken, ".pdf", "English").blocked
    monkeypatch.setenv("QYUNSLATION_PREFLIGHT_PROBE", "0")
    result = probe_document(broken, ".pdf", "English")
    assert result.issues == [] and not result.blocked

ENGLISH = (
    "The sponsor shall ensure that the clinical study is conducted in accordance with the protocol. "
    "Participants will receive the study drug once daily for twelve weeks and be assessed at each visit. "
) * 4
CHINESE = "申办者应当确保临床研究按照方案开展，受试者每天服用一次研究药物，持续十二周，并在每次访视时接受评估。" * 6


def _codes(result):
    return {issue["code"] for issue in result.issues}


def _pdf(path, pages, *, text=ENGLISH, scanned=(), password=None, table=False):
    import pymupdf
    from PIL import Image

    doc = pymupdf.open()
    png = io.BytesIO()
    Image.new("RGB", (600, 800), (230, 230, 230)).save(png, format="PNG")
    for index in range(pages):
        page = doc.new_page()
        if index in scanned:
            page.insert_image(page.rect, stream=png.getvalue())
        else:
            page.insert_textbox(pymupdf.Rect(50, 50, 550, 580 if table else 780), text, fontsize=9)
            if table:  # ten columns of cells: wide enough to count as complex
                for column in range(10):
                    for row in range(5):
                        page.draw_rect(pymupdf.Rect(50 + column * 45, 600 + row * 30, 95 + column * 45, 630 + row * 30))
                        page.insert_text((55 + column * 45, 620 + row * 30), "1", fontsize=8)
    kwargs = {"encryption": pymupdf.PDF_ENCRYPT_AES_256, "user_pw": password, "owner_pw": password} if password else {}
    doc.save(str(path), **kwargs)
    doc.close()
    return path


# --- PDF ----------------------------------------------------------------------------------------------------


def test_a_clean_pdf_has_no_issues_and_reports_its_pages(tmp_path):
    result = probe_document(_pdf(tmp_path / "a.pdf", 3), ".pdf", "English")
    assert result.issues == [] and not result.blocked
    assert result.manifest_summary["pages"] == 3 and result.manifest_summary["scan_detected"] is False


def test_an_encrypted_pdf_blocks_with_the_reason_the_impact_and_what_to_do(tmp_path):
    result = probe_document(_pdf(tmp_path / "a.pdf", 1, password="secret"), ".pdf", "English")
    assert result.blocked
    issue = result.issues[0]
    assert issue["code"] == "PDF_ENCRYPTED" and issue["severity"] == "blocker"
    assert issue["message"] and issue["impact"] and issue["action"]


def test_a_mostly_scanned_pdf_warns_with_a_count(tmp_path):
    result = probe_document(_pdf(tmp_path / "a.pdf", 10, scanned=set(range(8))), ".pdf", "English")
    assert "SCANNED_PDF" in _codes(result) and not result.blocked
    assert result.manifest_summary["scan_detected"] is True and result.manifest_summary["scanned_pages"] == 8
    assert "8/10" in next(i for i in result.issues if i["code"] == "SCANNED_PDF")["message"]


def test_a_few_scanned_pages_are_reported_without_calling_the_whole_file_a_scan(tmp_path):
    result = probe_document(_pdf(tmp_path / "a.pdf", 10, scanned={2, 3}), ".pdf", "English")
    assert "SCANNED_PAGES" in _codes(result) and "SCANNED_PDF" not in _codes(result)
    assert result.manifest_summary["scan_detected"] is False


def test_wide_tables_are_flagged(tmp_path):
    result = probe_document(_pdf(tmp_path / "a.pdf", 1, table=True), ".pdf", "English")
    assert "COMPLEX_TABLES" in _codes(result) and not result.blocked
    assert result.manifest_summary["tables"] == 1


def test_the_wrong_source_language_is_called_out_both_ways(tmp_path):
    chinese = probe_document(_pdf(tmp_path / "zh.pdf", 1, text="x"), ".pdf", "English")
    assert "LANGUAGE_MISMATCH" not in _codes(chinese)  # almost no text: nothing to judge
    zh_pdf = tmp_path / "zh2.txt"
    zh_pdf.write_text(CHINESE, encoding="utf-8")
    assert "LANGUAGE_MISMATCH" in _codes(probe_document(zh_pdf, ".txt", "English"))
    assert "LANGUAGE_MISMATCH" not in _codes(probe_document(zh_pdf, ".txt", "简体中文"))
    en = tmp_path / "en.txt"
    en.write_text(ENGLISH, encoding="utf-8")
    mismatch = probe_document(en, ".txt", "简体中文")
    assert "LANGUAGE_MISMATCH" in _codes(mismatch)
    assert "LANGUAGE_MISMATCH" not in _codes(probe_document(en, ".txt", "English"))


def test_a_corrupt_or_empty_file_is_a_blocker_not_a_crash(tmp_path):
    broken = tmp_path / "x.pdf"
    broken.write_bytes(b"this is not a pdf")
    assert probe_document(broken, ".pdf", "English").blocked
    empty = tmp_path / "e.txt"
    empty.write_text("   \n", encoding="utf-8")
    assert "TEXT_EMPTY" in _codes(probe_document(empty, ".txt", "English"))


# --- Office and images ----------------------------------------------------------------------------------------


def test_docx_reports_tables_and_flags_nested_or_wide_ones(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph(ENGLISH)
    document.add_table(rows=2, cols=3)
    path = tmp_path / "a.docx"
    document.save(str(path))
    plain = probe_document(path, ".docx", "English")
    assert plain.manifest_summary["tables"] == 1 and "TABLES" in _codes(plain) and not plain.blocked

    wide = docx.Document()
    wide.add_table(rows=2, cols=9)
    wide_path = tmp_path / "w.docx"
    wide.save(str(wide_path))
    assert "COMPLEX_TABLES" in _codes(probe_document(wide_path, ".docx", "English"))


def test_a_password_protected_office_file_is_recognised_by_its_container(tmp_path):
    path = tmp_path / "locked.docx"
    path.write_bytes(probe_module.OLE_MAGIC + b"\x00" * 200)
    result = probe_document(path, ".docx", "English")
    assert result.blocked and result.issues[0]["code"] == "DOCX_ENCRYPTED"
    junk = tmp_path / "junk.pptx"
    junk.write_bytes(b"nope")
    assert probe_document(junk, ".pptx", "English").issues[0]["code"] == "PPTX_UNREADABLE"
    bad_zip = tmp_path / "bad.docx"
    with zipfile.ZipFile(bad_zip, "w") as archive:
        archive.writestr("hello.txt", "not a word document")
    assert probe_document(bad_zip, ".docx", "English").blocked


def test_pptx_counts_slides_and_flags_an_empty_deck(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Study overview"
    path = tmp_path / "a.pptx"
    deck.save(str(path))
    assert probe_document(path, ".pptx", "English").manifest_summary["pages"] == 1
    empty = tmp_path / "e.pptx"
    Presentation().save(str(empty))
    assert probe_document(empty, ".pptx", "English").blocked


def test_a_small_image_warns_about_ocr_quality(tmp_path):
    from PIL import Image

    small = tmp_path / "s.png"
    Image.new("RGB", (400, 300), "white").save(small)
    assert "LOW_RESOLUTION" in _codes(probe_document(small, ".png", "English"))
    large = tmp_path / "l.png"
    Image.new("RGB", (1600, 1200), "white").save(large)
    assert "LOW_RESOLUTION" not in _codes(probe_document(large, ".png", "English"))


def test_the_probe_never_raises_and_is_time_boxed(tmp_path, monkeypatch):
    path = _pdf(tmp_path / "a.pdf", 3)

    def boom(*args, **kwargs):
        raise RuntimeError("library blew up")

    real_probe = probe_module._probe_pdf
    monkeypatch.setattr(probe_module, "_probe_pdf", boom)
    result = probe_document(path, ".pdf", "English")
    assert [issue["code"] for issue in result.issues] == ["PROBE_FAILED"] and not result.blocked
    monkeypatch.setattr(probe_module, "_probe_pdf", real_probe)
    monkeypatch.setattr(probe_module, "PROBE_BUDGET_SECONDS", 0.0)
    assert probe_document(path, ".pdf", "English").manifest_summary["pages"] == 3  # stops sampling, still answers


# --- through the API --------------------------------------------------------------------------------------------


@pytest.fixture()
def client(monkeypatch, tmp_path):
    reset_engine()
    monkeypatch.setenv("QYUNSLATION_DATABASE_URL", "sqlite+pysqlite:///:memory:")
    monkeypatch.setenv("QYUNSLATION_DEV_AUTH_BYPASS", "1")
    monkeypatch.setenv("QYUNSLATION_ENV", "development")
    monkeypatch.setenv("QYUNSLATION_PREFLIGHT_ROOT", str(tmp_path / "preflights"))
    monkeypatch.setenv("QYUNSLATION_UPLOAD_ROOT", str(tmp_path / "uploads"))
    Base.metadata.create_all(init_engine("sqlite+pysqlite:///:memory:"))
    app = FastAPI()
    app.include_router(api_v1_router)
    with TestClient(app) as current:
        yield current
    reset_engine()


def _upload(client, path, name):
    return client.post(
        "/api/v1/preflights", headers=OWNER,
        data={"source_language": "English", "target_language": "简体中文"},
        files={"file": (name, path.read_bytes(), "application/octet-stream")},
    )


def test_an_encrypted_pdf_is_blocked_at_upload_and_cannot_start_a_run(client, tmp_path):
    response = _upload(client, _pdf(tmp_path / "locked.pdf", 1, password="pw"), "locked.pdf")
    assert response.status_code in {200, 201}
    body = response.json()
    assert body["state"] == "blocked" and body["issues"][0]["code"] == "PDF_ENCRYPTED"
    created = client.post(
        "/api/v1/translation-runs", headers={**OWNER, "Idempotency-Key": "k1"}, json={"preflight_id": body["id"]}
    )
    assert created.status_code in {409, 422}


def test_a_scanned_pdf_carries_its_limits_in_the_preflight(client, tmp_path):
    body = _upload(client, _pdf(tmp_path / "scan.pdf", 6, scanned=set(range(6))), "scan.pdf").json()
    assert body["state"] == "ready"
    assert body["manifest_summary"]["pages"] == 6 and body["manifest_summary"]["scan_detected"] is True
    assert any(issue["code"] == "SCANNED_PDF" for issue in body["issues"])


def test_the_resumable_upload_path_runs_the_probe_too(client, tmp_path):
    path = _pdf(tmp_path / "locked.pdf", 1, password="pw")
    data = path.read_bytes()
    created = client.post("/api/v1/upload-sessions", headers=OWNER, json={"filename": "locked.pdf", "total_size": len(data)}).json()
    upload_id = (created.get("upload_session") or created)["id"]
    client.patch(
        f"/api/v1/upload-sessions/{upload_id}",
        headers={**OWNER, "Content-Range": f"bytes 0-{len(data) - 1}/{len(data)}", "Content-Type": "application/octet-stream"},
        content=data,
    )
    done = client.post(
        f"/api/v1/upload-sessions/{upload_id}/complete", headers=OWNER,
        json={"source_language": "English", "target_language": "简体中文", "profile": "临床研究文档"},
    )
    assert done.status_code in {200, 201}, done.text
    assert done.json()["state"] == "blocked"


def test_the_server_publishes_the_formats_and_size_limit_the_upload_box_must_use(client):
    upload = client.get("/api/v1/settings/schema", headers=OWNER).json()["upload"]
    assert ".pdf" in upload["allowed_extensions"] and ".docx" in upload["allowed_extensions"]
    assert upload["max_bytes"] > 0
