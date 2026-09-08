"""Generate small, deterministic PLAN-030 structure fixtures.

The generated binaries are intentionally not committed. Their SHA-256 values are
recorded in ``catalog.v1.json`` so generation drift fails loudly.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.shared import Inches
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as PptxInches


FIXED_TIME = datetime(2025, 1, 1, tzinfo=UTC)
ZIP_TIME = (2025, 1, 1, 0, 0, 0)


def _pdf_bytes(*, columns: int) -> bytes:
    if columns not in {1, 2}:
        raise ValueError("columns must be 1 or 2")

    title = "Single-column research article" if columns == 1 else "Double-column review"
    text_positions = (
        [(72, 700, "Body text follows reading order across one column.")]
        if columns == 1
        else [
            (72, 700, "Left column body text and citation 1."),
            (324, 700, "Right column body text and citation 2."),
        ]
    )
    commands = [
        "BT /F1 18 Tf 72 752 Td (" + title + ") Tj ET",
        *[
            f"BT /F1 10 Tf {x} {y} Td ({text}) Tj ET"
            for x, y, text in text_positions
        ],
        "0.7 w 72 420 468 190 re S",
        "BT /F1 10 Tf 72 404 Td (Figure 1. Cross-column result.) Tj ET",
        "72 250 468 120 re S",
        "72 290 m 540 290 l S",
        "306 250 m 306 370 l S",
        "BT /F1 10 Tf 72 234 Td (Table 1. Structured outcomes.) Tj ET",
    ]
    stream = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(output)


def _assemble_pdf(objects: list[bytes]) -> bytes:
    """把已编号的对象体拼成最小 PDF。对象 1 必须是 Catalog。"""
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(output)


def _stream_object(payload: bytes) -> bytes:
    return (
        b"<< /Length "
        + str(len(payload)).encode("ascii")
        + b" >>\nstream\n"
        + payload
        + b"\nendstream"
    )


# 与仓外扫描件金样等价的合成件：整页图 + 可选 OCR 文字层。
# 真实件是 20 页 FDA PIND，不入库；这里复现的是形态与页数，不是内容。
SCANNED_EQUIVALENT_PAGES = 20
# 与仓外幻灯金样等价：16:9、无题注、每页一个矢量区域。
SLIDE_EQUIVALENT_PAGES = 12
_SLIDE_DRAWINGS_PER_PAGE = 8


def _scanned_pdf_bytes(*, pages: int, text_layer: bool) -> bytes:
    """扫描态 PDF：每页一张整页灰度图。

    `text_layer=False` 让 `page_representation()` 判 SCANNED（文字 < 20 字符）；
    `True` 追加足够长的文字层，判 HYBRID（有文字又压着整页图），复现 hpd-ocr 产物形态。
    """
    width, height = 612, 792
    # 8x8 DeviceGray 原始位图，够小且完全确定
    pixels = bytes(range(0, 256, 4))[:64]
    image_object = (
        b"<< /Type /XObject /Subtype /Image /Width 8 /Height 8 "
        b"/ColorSpace /DeviceGray /BitsPerComponent 8 /Length "
        + str(len(pixels)).encode("ascii")
        + b" >>\nstream\n"
        + pixels
        + b"\nendstream"
    )

    first_page_number = 5
    page_refs = " ".join(
        f"{first_page_number + i} 0 R" for i in range(pages)
    )
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        (
            b"<< /Type /Pages /Kids ["
            + page_refs.encode("ascii")
            + b"] /Count "
            + str(pages).encode("ascii")
            + b" >>"
        ),
        image_object,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    contents_start = first_page_number + pages
    for index in range(pages):
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
                f"/Resources << /XObject << /Im0 3 0 R >> /Font << /F1 4 0 R >> >> "
                f"/Contents {contents_start + index} 0 R >>"
            ).encode("ascii")
        )

    for index in range(pages):
        commands = [f"q {width} 0 0 {height} 0 0 cm /Im0 Do Q"]
        if text_layer:
            commands.append(
                "BT /F1 10 Tf 72 700 Td "
                f"(Recovered text layer for scanned page {index + 1} of {pages}.) Tj ET"
            )
        objects.append(_stream_object("\n".join(commands).encode("ascii")))

    return _assemble_pdf(objects)


def _slide_pdf_bytes(*, pages: int) -> bytes:
    """幻灯导出态 PDF：16:9 画布、无 Figure/Table 题注、每页一个矢量区域。

    区域由纵向邻近的矩形聚类而成，因此每页画 `_SLIDE_DRAWINGS_PER_PAGE` 个
    彼此靠近的小矩形，超过 `SLIDE_MIN_DRAWINGS` 门槛且聚成一簇。
    文字保持短句，避免形成 `_long_text_rects` 的防触碰块。
    """
    width, height = 720, 405  # 16:9，宽高比 1.78 ≥ 1.55

    first_page_number = 4
    page_refs = " ".join(f"{first_page_number + i} 0 R" for i in range(pages))
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        (
            b"<< /Type /Pages /Kids ["
            + page_refs.encode("ascii")
            + b"] /Count "
            + str(pages).encode("ascii")
            + b" >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    contents_start = first_page_number + pages
    for index in range(pages):
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
                f"/Resources << /Font << /F1 3 0 R >> >> "
                f"/Contents {contents_start + index} 0 R >>"
            ).encode("ascii")
        )

    for index in range(pages):
        # 标题与副标题都放在顶部，与下方矢量区域不重叠；每块都短于
        # `_long_text_rects` 的 60 字符门槛，不会形成防触碰块
        commands = [
            f"BT /F1 16 Tf 48 {height - 48} Td (Slide {index + 1} overview) Tj ET",
            f"BT /F1 11 Tf 48 {height - 72} Td (Study design and endpoints) Tj ET",
            "0.8 w",
        ]
        # 一簇纵向相邻的矩形：聚类后成为单个可译区。行距必须让
        # `find_safe_vector_figures` 的纵向邻近判据成立（并簇后 y 跨度 < 40）
        for row in range(_SLIDE_DRAWINGS_PER_PAGE):
            y = 90 + row * 20
            commands.append(f"120 {y} 200 18 re S")
        objects.append(_stream_object("\n".join(commands).encode("ascii")))

    return _assemble_pdf(objects)


# --- PLAN-030h H4：栏式版式金样 ------------------------------------------------
#
# 全部用裸内容流手写，不走 pymupdf 的文档生成：pymupdf 的 tobytes(garbage=4)
# 会重排 xref，产物字节受进程内累积状态影响，在全仓并发跑门时出现过同一夹具
# 两次生成 hash 不同、catalog 复现断言转红。裸流没有这个问题。
#
# 排版上有两处约束必须守住，否则测不出真实的栏式行为：
#   - 同一基线 y 上的多列文字会被 PyMuPDF 并成一个横跨整页的块，narrow 块数
#     不足，栏式判定看不到栏。各列整体错开 COLUMN_BASELINE_STAGGER 解决。
#   - 行距超过约 34pt 时列内各行不再聚成一块，碎成一行一块。保持 LINE_GAP。

MULTI_COLUMN_LAYOUTS: dict[str, tuple[float, ...]] = {
    "three": (60.0, 232.0, 404.0),
    "four": (50.0, 190.0, 330.0, 470.0),
}

COLUMN_BASELINE_STAGGER = 3.0
COLUMN_LINE_GAP = 11.0
COLUMN_LINES = 8


def _text_command(x: float, y: float, size: float, text: str) -> str:
    return f"BT /F1 {size:g} Tf {x:.1f} {y:.1f} Td ({text}) Tj ET"


def _column_commands(
    column_xs: tuple[float, ...],
    *,
    top: float,
    label: str,
    size: float = 8.0,
    line_gap: float = COLUMN_LINE_GAP,
    lines: int = COLUMN_LINES,
    stagger: float = COLUMN_BASELINE_STAGGER,
) -> list[str]:
    commands = []
    for index, x in enumerate(column_xs):
        for line in range(lines):
            y = top - line * line_gap - index * stagger
            commands.append(
                _text_command(
                    x,
                    y,
                    size,
                    f"{label} column {index} line {line} of running body text",
                )
            )
    return commands


def _single_page_pdf(commands: list[str], *, width: float, height: float) -> bytes:
    stream = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width:.0f} {height:.0f}] "
            "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ).encode("ascii"),
        _stream_object(stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    return _assemble_pdf(objects)


def _multi_column_pdf_bytes(kind: str) -> bytes:
    commands = [_text_command(72.0, 740.0, 16.0, f"{kind.capitalize()} layout gold sample")]
    commands += _column_commands(MULTI_COLUMN_LAYOUTS[kind], top=700.0, label=kind.capitalize())
    return _single_page_pdf(commands, width=612.0, height=792.0)


def _mixed_columns_pdf_bytes() -> bytes:
    """三页文档，逐页栏式不同：单栏扉页、双栏正文、三栏附录。

    栏式是逐页判定的，「混合栏」在这个模型里只能是文档级的观察结果。
    """
    pages: list[list[str]] = []

    front = [_text_command(72.0, 740.0, 18.0, "Mixed layout gold sample")]
    for line in range(10):
        front.append(
            _text_command(
                72.0,
                700.0 - line * 13.0,
                9.0,
                f"Front matter line {line} runs the full measure of the page width",
            )
        )
    pages.append(front)

    body = [_text_command(72.0, 740.0, 16.0, "Two-column body")]
    body += _column_commands((72.0, 330.0), top=700.0, label="Body")
    pages.append(body)

    appendix = [_text_command(72.0, 740.0, 16.0, "Three-column appendix")]
    appendix += _column_commands(MULTI_COLUMN_LAYOUTS["three"], top=700.0, label="Appendix")
    pages.append(appendix)

    return _multi_page_pdf(pages, width=612.0, height=792.0)


def _multi_page_pdf(pages: list[list[str]], *, width: float, height: float) -> bytes:
    page_count = len(pages)
    # 对象布局：1 Catalog、2 Pages、3..(2+n) Page、随后各页 Contents，最后 Font
    first_page_xref = 3
    first_content_xref = first_page_xref + page_count
    font_xref = first_content_xref + page_count

    kids = " ".join(f"{first_page_xref + i} 0 R" for i in range(page_count))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode("ascii"),
    ]
    for index in range(page_count):
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width:.0f} {height:.0f}] "
                f"/Resources << /Font << /F1 {font_xref} 0 R >> >> "
                f"/Contents {first_content_xref + index} 0 R >>"
            ).encode("ascii")
        )
    for commands in pages:
        objects.append(_stream_object("\n".join(commands).encode("ascii")))
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    return _assemble_pdf(objects)


POSTER_SIZE = (3370.0, 2384.0)
POSTER_BAND_TOPS = (2000.0, 1240.0, 480.0)
POSTER_COLUMN_LEFTS = (160.0, 1180.0, 2200.0)


def _poster_pdf_bytes() -> bytes:
    """A0 横版海报：三条分区带，每带三块，标题横跨全宽。"""
    width, height = POSTER_SIZE
    commands = [_text_command(160.0, 2240.0, 72.0, "Poster section gold sample")]
    for band, top in enumerate(POSTER_BAND_TOPS):
        for column, left in enumerate(POSTER_COLUMN_LEFTS):
            commands.append(f"2 w {left:.0f} {top - 660:.0f} 940 660 re S")
            for line in range(6):
                commands.append(
                    _text_command(
                        left + 30.0,
                        top - 60.0 - line * 30.0 - column * 22.0,
                        24.0,
                        f"Band {band} panel {column} line {line} describing results.",
                    )
                )
    return _single_page_pdf(commands, width=width, height=height)


# --- PLAN-030h H3：同一逻辑内容的四种承载物 ----------------------------------
#
# 四份夹具承载**同一份内容**：一段正文、一张带 Figure 1 题注的图、一张带
# Table 1 题注的表。跨格式对账断言据此比对语义对象集合，差异必须由
# content_profile / container_mode / output_editability 解释。

PARITY_BODY = "Body text follows reading order across one column."
PARITY_FIGURE_CAPTION = "Figure 1. Cross-format parity figure."
PARITY_TABLE_CAPTION = "Table 1. Cross-format parity table."
PARITY_TITLE = "Cross-format parity fixture"


def _parity_pdf_bytes() -> bytes:
    commands = [
        f"BT /F1 18 Tf 72 752 Td ({PARITY_TITLE}) Tj ET",
        f"BT /F1 10 Tf 72 700 Td ({PARITY_BODY}) Tj ET",
        "0.7 w 72 420 468 190 re S",
        f"BT /F1 10 Tf 72 404 Td ({PARITY_FIGURE_CAPTION}) Tj ET",
        "72 250 468 120 re S",
        "72 290 m 540 290 l S",
        "306 250 m 306 370 l S",
        f"BT /F1 10 Tf 72 234 Td ({PARITY_TABLE_CAPTION}) Tj ET",
    ]
    stream = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        _stream_object(stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    return _assemble_pdf(objects)


def _parity_docx_bytes() -> bytes:
    document = Document()
    document.core_properties.title = PARITY_TITLE
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME

    document.add_heading(PARITY_TITLE, 0)
    document.add_paragraph(PARITY_BODY)
    document.add_picture(BytesIO(_png_bytes()), width=Inches(2.2))
    document.add_paragraph(PARITY_FIGURE_CAPTION)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Outcome"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Stable"
    table.cell(1, 1).text = "Yes"
    document.add_paragraph(PARITY_TABLE_CAPTION)

    stream = BytesIO()
    document.save(stream)
    return _canonicalize_ooxml(stream.getvalue())


def _parity_pptx_bytes() -> bytes:
    presentation = Presentation()
    presentation.core_properties.title = PARITY_TITLE
    presentation.core_properties.created = FIXED_TIME
    presentation.core_properties.modified = FIXED_TIME
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    title = slide.shapes.add_textbox(
        PptxInches(0.5), PptxInches(0.3), PptxInches(9), PptxInches(0.6)
    )
    title.text_frame.text = PARITY_TITLE
    body = slide.shapes.add_textbox(
        PptxInches(0.5), PptxInches(1.0), PptxInches(9), PptxInches(0.5)
    )
    body.text_frame.text = PARITY_BODY

    slide.shapes.add_picture(
        BytesIO(_png_bytes()), PptxInches(0.5), PptxInches(1.7), width=PptxInches(4)
    )
    figure_caption = slide.shapes.add_textbox(
        PptxInches(0.5), PptxInches(4.2), PptxInches(4), PptxInches(0.4)
    )
    figure_caption.text_frame.text = PARITY_FIGURE_CAPTION

    table = slide.shapes.add_table(
        2, 2, PptxInches(5.0), PptxInches(1.7), PptxInches(4), PptxInches(1.6)
    ).table
    table.cell(0, 0).text = "Outcome"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Stable"
    table.cell(1, 1).text = "Yes"
    table_caption = slide.shapes.add_textbox(
        PptxInches(5.0), PptxInches(4.2), PptxInches(4), PptxInches(0.4)
    )
    table_caption.text_frame.text = PARITY_TABLE_CAPTION

    stream = BytesIO()
    presentation.save(stream)
    return _canonicalize_ooxml(stream.getvalue())


def _parity_png_bytes() -> bytes:
    """同一内容的整页栅格承载：结构在这里被压平，只剩像素。"""
    image = Image.new("RGB", (612, 792), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    draw.text((72, 40), PARITY_TITLE, fill=(0, 0, 0))
    draw.text((72, 92), PARITY_BODY, fill=(0, 0, 0))
    draw.rectangle((72, 182, 540, 372), outline=(0, 0, 0), width=1)
    draw.text((72, 388), PARITY_FIGURE_CAPTION, fill=(0, 0, 0))
    draw.rectangle((72, 422, 540, 542), outline=(0, 0, 0), width=1)
    draw.line((72, 482, 540, 482), fill=(0, 0, 0), width=1)
    draw.line((306, 422, 306, 542), fill=(0, 0, 0), width=1)
    draw.text((72, 558), PARITY_TABLE_CAPTION, fill=(0, 0, 0))
    return _save_image(image, "PNG", optimize=False)


def _base_image(*, poster: bool = False, alpha: bool = False) -> Image.Image:
    size = (900, 1200) if poster else (320, 180)
    mode = "RGBA" if alpha else "RGB"
    background = (246, 248, 252, 0 if alpha else 255) if alpha else (246, 248, 252)
    image = Image.new(mode, size, background)
    draw = ImageDraw.Draw(image)
    width, height = size
    draw.rectangle((16, 16, width - 16, height - 16), outline=(11, 95, 130), width=4)
    draw.rectangle((32, 48, width // 2 - 8, height - 48), fill=(218, 238, 245))
    draw.rectangle((width // 2 + 8, 48, width - 32, height - 48), fill=(231, 239, 208))
    draw.text((32, 24), "PLAN-030 TRANSLATION FIXTURE", fill=(18, 42, 66))
    draw.text((48, 70), "Figure 1", fill=(18, 42, 66))
    draw.text((width // 2 + 24, 70), "Table 1", fill=(18, 42, 66))
    return image


def _save_image(image: Image.Image, format_name: str, **kwargs: object) -> bytes:
    stream = BytesIO()
    image.save(stream, format=format_name, **kwargs)
    return stream.getvalue()


def _png_bytes() -> bytes:
    return _save_image(_base_image(poster=True, alpha=True), "PNG", optimize=False)


def _jpeg_bytes() -> bytes:
    exif = Image.Exif()
    exif[274] = 6
    return _save_image(
        _base_image(),
        "JPEG",
        quality=85,
        subsampling=0,
        optimize=False,
        progressive=False,
        exif=exif,
    )


def _webp_bytes() -> bytes:
    return _save_image(_base_image(), "WEBP", lossless=True, method=6, exact=True)


def _bmp_bytes() -> bytes:
    return _save_image(_base_image(), "BMP")


def _tiff_bytes() -> bytes:
    first = _base_image()
    second = _base_image().transpose(Image.Transpose.FLIP_LEFT_RIGHT)
    stream = BytesIO()
    first.save(
        stream,
        format="TIFF",
        save_all=True,
        append_images=[second],
        compression="raw",
    )
    return stream.getvalue()


def _canonicalize_ooxml(data: bytes) -> bytes:
    source = BytesIO(data)
    output = BytesIO()
    with zipfile.ZipFile(source, "r") as input_zip, zipfile.ZipFile(
        output,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
        strict_timestamps=True,
    ) as output_zip:
        for name in sorted(input_zip.namelist()):
            info = zipfile.ZipInfo(name, ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            output_zip.writestr(info, input_zip.read(name))
    return output.getvalue()


def _docx_bytes() -> bytes:
    document = Document()
    document.core_properties.title = "PLAN-030 deterministic review fixture"
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME

    header = document.sections[0].header
    header.paragraphs[0].text = "Deterministic header for review"
    footer = document.sections[0].footer
    footer.paragraphs[0].text = "Deterministic footer for review"

    section_properties = document.sections[0]._sectPr
    columns = section_properties.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        section_properties.append(columns)
    columns.set(qn("w:num"), "2")
    columns.set(qn("w:space"), "720")

    document.add_heading("Deterministic review", 0)
    document.add_paragraph("Two-column narrative with a native table and shared image part.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Outcome"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Stable"
    table.cell(1, 1).text = "Yes"

    image = _png_bytes()
    document.add_picture(BytesIO(image), width=Inches(2.2))
    document.add_paragraph("Figure 1. First occurrence.")
    document.add_picture(BytesIO(image), width=Inches(1.6))
    document.add_paragraph("Figure 1. Reused occurrence.")

    textbox_para = document.add_paragraph()
    run = textbox_para.add_run()
    txbx = OxmlElement("w:txbxContent")
    inner_p = OxmlElement("w:p")
    inner_r = OxmlElement("w:r")
    inner_t = OxmlElement("w:t")
    inner_t.text = "Standalone text box content."
    inner_r.append(inner_t)
    inner_p.append(inner_r)
    txbx.append(inner_p)
    run._r.append(txbx)

    document.add_page_break()
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Second section body in single flow.")

    stream = BytesIO()
    document.save(stream)
    return _canonicalize_ooxml(stream.getvalue())


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    presentation.core_properties.title = "PLAN-030 deterministic presentation fixture"
    presentation.core_properties.created = FIXED_TIME
    presentation.core_properties.modified = FIXED_TIME
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    title = slide.shapes.add_textbox(
        PptxInches(0.5), PptxInches(0.3), PptxInches(9), PptxInches(0.7)
    )
    title.text_frame.text = "Native slide text"
    table = slide.shapes.add_table(
        2, 2, PptxInches(0.5), PptxInches(1.3), PptxInches(4), PptxInches(2)
    ).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Coverage"
    table.cell(1, 1).text = "100%"
    slide.shapes.add_picture(
        BytesIO(_png_bytes()),
        PptxInches(5),
        PptxInches(1.3),
        width=PptxInches(4),
    )

    stream = BytesIO()
    presentation.save(stream)
    return _canonicalize_ooxml(stream.getvalue())


def _caption_span_gap_pdf_bytes() -> bytes:
    """复现 ScienceDirect 题注：同一行两个 span，中间没有空格字形。

    旧拼接把 'Table 2' + 'Response...' 合成 'Table 2Response...'，正则因此丢表。
    正文里的 'see Table 4' 不得被计成对象。
    """
    commands = [
        _text_command(72.0, 740.0, 16.0, "Caption span-gap gold sample"),
        _text_command(72.0, 700.0, 10.0, "Figure 1. Study flow."),
        # 两个 Tj 紧挨着，中间不写空格字符
        "BT /F1 10 Tf 72.0 640.0 Td (Table 2) Tj 42 0 Td (Response to Dupilumab in Acute GvHD.) Tj ET",
        "BT /F1 10 Tf 72.0 600.0 Td (Table 3) Tj 42 0 Td (Response in Chronic GvHD.) Tj ET",
        "BT /F1 10 Tf 72.0 560.0 Td (Table 4) Tj 42 0 Td (Baseline and Transplant Characteristics.) Tj ET",
        _text_command(72.0, 480.0, 10.0, "Patients were enrolled, see Table 4 for details."),
    ]
    return _single_page_pdf(commands, width=612.0, height=792.0)


def _landscape_frame_table_pdf_bytes() -> bytes:
    """复现侧放整页框线表：题注在框左侧，上下横线间距超过 ROW_GAP_BREAK。

    旧算法只向下找横线群，两条框线相距 >0.20 页高就被拆掉。033b 用两横+两竖
    封闭框回退圈定。
    """
    commands = [
        "0 0 0 rg",
        "94 54 234 0.5 re f",
        "94 740 234 0.5 re f",
        "94 54 0.5 686 re f",
        "328 54 0.5 686 re f",
        _text_command(52.0, 80.0, 10.0, "Table 1"),
        _text_command(110.0, 700.0, 8.0, "Age (y)"),
        _text_command(180.0, 700.0, 8.0, "12"),
        _text_command(110.0, 620.0, 8.0, "Sex"),
        _text_command(180.0, 620.0, 8.0, "M"),
        _text_command(72.0, 20.0, 9.0, "Footer outside the framed table stays body text."),
    ]
    return _single_page_pdf(commands, width=612.0, height=792.0)


GENERATORS: dict[str, Callable[[], bytes]] = {
    "single-column.pdf": lambda: _pdf_bytes(columns=1),
    "double-column.pdf": lambda: _pdf_bytes(columns=2),
    "scanned-equivalent.pdf": lambda: _scanned_pdf_bytes(
        pages=SCANNED_EQUIVALENT_PAGES, text_layer=False
    ),
    "scanned-equivalent.hpd-ocr.pdf": lambda: _scanned_pdf_bytes(
        pages=SCANNED_EQUIVALENT_PAGES, text_layer=True
    ),
    "slide-equivalent.pdf": lambda: _slide_pdf_bytes(pages=SLIDE_EQUIVALENT_PAGES),
    "review.docx": _docx_bytes,
    "poster.png": _png_bytes,
    "photo.jpg": _jpeg_bytes,
    "diagram.webp": _webp_bytes,
    "scan.bmp": _bmp_bytes,
    "multipage.tiff": _tiff_bytes,
    "presentation.pptx": _pptx_bytes,
    "parity.pdf": _parity_pdf_bytes,
    "parity.docx": _parity_docx_bytes,
    "parity.pptx": _parity_pptx_bytes,
    "parity.png": _parity_png_bytes,
    "three-column.pdf": lambda: _multi_column_pdf_bytes("three"),
    "four-column.pdf": lambda: _multi_column_pdf_bytes("four"),
    "mixed-columns.pdf": _mixed_columns_pdf_bytes,
    "poster-sections.pdf": _poster_pdf_bytes,
    "caption-span-gap.pdf": _caption_span_gap_pdf_bytes,
    "landscape-frame-table.pdf": _landscape_frame_table_pdf_bytes,
}


def generate_all(output_dir: Path) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    hashes: dict[str, str] = {}
    for relative_path, generator in GENERATORS.items():
        data = generator()
        target = output_dir / relative_path
        target.write_bytes(data)
        hashes[relative_path] = hashlib.sha256(data).hexdigest()
    return hashes


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(generate_all(args.output), indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
