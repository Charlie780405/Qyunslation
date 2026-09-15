# SPDX-License-Identifier: MPL-2.0
"""Server-side client used by the patched pdf2zh Gradio callback.

The browser never sees this client, its HMAC secret, or the internal sidecar
route.  Errors intentionally collapse to ``WorkbenchTermBridgeUnavailable`` so
the translation can complete with a truthful degraded term-review state.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import tempfile
import time
import uuid
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import httpx

from qyunslation.workbench.evidence import BilingualTermEvidence
from qyunslation.workbench.security import BRIDGE_SECRET_ENV, sign_bridge_request
from qyunslation.structure.babeldoc_policy import ReferencePreserveGate
from qyunslation.structure.references import text_excluding_reference_entries

BRIDGE_URL_ENV = "QYUNSLATION_WORKBENCH_BRIDGE_URL"
_DEFAULT_BRIDGE_URL = "http://127.0.0.1:8010/internal/workbench/v1"
_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".html", ".htm", ".csv", ".json"}
_IMAGE_SUFFIXES = {".bmp", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"}
_MEANINGFUL_TEXT = re.compile(r"[A-Za-z\u4e00-\u9fff]")


def _ocr_policy_asset_limit() -> int:
    try:
        return max(0, int(os.environ.get("QYUNSLATION_TERM_POLICY_OCR_IMAGE_LIMIT", "12")))
    except ValueError:
        return 12


def _cell_term_pair(source_text: str, target_text: str) -> tuple[str | None, str | None]:
    """Use an Office cell only when it is short enough to be a reviewable unit."""
    source = " ".join((source_text or "").split())
    target = " ".join((target_text or "").split())
    if not target or not (2 <= len(source) <= 160) or not _MEANINGFUL_TEXT.search(source):
        return None, None
    return source, target[:512]


class WorkbenchTermBridgeUnavailable(RuntimeError):
    """The translator remains available, but term review cannot be trusted."""

    def __init__(
        self,
        message: str,
        *,
        status_code: int | None = None,
        detail: str | None = None,
    ):
        super().__init__(message)
        self.status_code = status_code
        self.detail = detail


def _bridge_base_url() -> str:
    return (os.environ.get(BRIDGE_URL_ENV) or _DEFAULT_BRIDGE_URL).rstrip("/")


def _bridge_secret() -> str:
    secret = (os.environ.get(BRIDGE_SECRET_ENV) or "").strip()
    if not secret:
        raise WorkbenchTermBridgeUnavailable("专业词库暂不可用")
    return secret


def _language(value: str | None) -> str:
    normalized = (value or "").strip().casefold()
    if normalized in {"english", "en", "英文", "英语"}:
        return "en"
    if normalized in {"simplified chinese", "chinese", "zh", "中文", "简体中文"}:
        return "zh"
    return normalized or "auto"


def _ocr_text(path: Path) -> str:
    """Return stable OCR reading order without ever blocking translation setup."""
    try:
        from qyunslation.extensions.image_translate import ocr_image

        return "\n".join(
            str(box[4]).strip()
            for box in ocr_image(path)
            if len(box) >= 5 and str(box[4]).strip()
        )
    except Exception:
        return ""


def _ocr_blob(blob: bytes, suffix: str) -> str:
    if not blob or _ocr_policy_asset_limit() <= 0:
        return ""
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=True) as handle:
        handle.write(blob)
        handle.flush()
        return _ocr_text(Path(handle.name))


def _package_media_ocr_text(path: Path, prefix: str) -> list[str]:
    """Read a bounded number of Office image assets for pre-translation terms."""
    limit = _ocr_policy_asset_limit()
    if limit <= 0:
        return []
    rows: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            media = sorted(
                member
                for member in archive.namelist()
                if member.startswith(prefix)
                and Path(member).suffix.casefold() in _IMAGE_SUFFIXES
            )[:limit]
            for member in media:
                text = _ocr_blob(archive.read(member), Path(member).suffix)
                if text:
                    rows.append(text)
    except (OSError, zipfile.BadZipFile):
        return []
    return rows


def _pdf_image_ocr_text(document: Any) -> list[str]:
    limit = _ocr_policy_asset_limit()
    if limit <= 0:
        return []
    rows: list[str] = []
    seen_xrefs: set[int] = set()
    for page in document:
        for image in page.get_images(full=True):
            xref = int(image[0])
            if xref in seen_xrefs:
                continue
            seen_xrefs.add(xref)
            try:
                extracted = document.extract_image(xref)
                text = _ocr_blob(
                    extracted.get("image", b""),
                    f".{str(extracted.get('ext') or 'png')}",
                )
            except Exception:
                continue
            if text:
                rows.append(text)
            if len(seen_xrefs) >= limit:
                return rows
    return rows


def _read_text(path: Path, *, limit: int = 1_500_000) -> str:
    try:
        if path.suffix.lower() in _TEXT_SUFFIXES:
            return path.read_text(encoding="utf-8", errors="replace")[:limit]
        if path.suffix.lower() in _IMAGE_SUFFIXES:
            return _ocr_text(path)[:limit]
        if path.suffix.lower() == ".pdf":
            import fitz

            with fitz.open(path) as document:
                parts = [text_excluding_reference_entries(document)]
                parts.extend(_pdf_image_ocr_text(document))
                return "\n".join(part for part in parts if part)[:limit]
        if path.suffix.lower() == ".docx":
            from docx import Document

            document = Document(path)
            gate = ReferencePreserveGate()
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text and not gate.should_preserve(paragraph.text)
            ]
            for table in document.tables:
                parts.extend(cell.text for row in table.rows for cell in row.cells)
            parts.extend(_package_media_ocr_text(path, "word/media/"))
            return "\n".join(parts)[:limit]
        if path.suffix.lower() in {".ppt", ".pptx"}:
            from pptx import Presentation

            presentation = Presentation(path)
            parts: list[str] = []
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        parts.append(shape.text)
                    if getattr(shape, "has_table", False):
                        parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            parts.extend(_package_media_ocr_text(path, "ppt/media/"))
            return "\n".join(parts)[:limit]
    except Exception:
        return ""
    return ""


def _bridge_request(method: str, path: str, body: dict[str, Any]) -> dict[str, Any]:
    raw = json.dumps(body, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    signed_path = path.partition("?")[0]
    headers = sign_bridge_request(
        method=method,
        path=signed_path,
        body=raw,
        secret=_bridge_secret(),
        nonce=uuid.uuid4().hex,
        timestamp=int(time.time()),
    )
    try:
        response = httpx.request(
            method,
            _bridge_base_url() + path.removeprefix("/internal/workbench/v1"),
            content=raw,
            headers=headers,
            timeout=httpx.Timeout(20.0, connect=3.0),
        )
        response.raise_for_status()
        payload = response.json()
    except httpx.HTTPStatusError as exc:
        detail = None
        try:
            body = exc.response.json()
            detail = body.get("detail") if isinstance(body, dict) else None
        except ValueError:
            detail = None
        raise WorkbenchTermBridgeUnavailable(
            "专业词库暂不可用",
            status_code=exc.response.status_code,
            detail=str(detail) if detail is not None else None,
        ) from exc
    except (httpx.HTTPError, ValueError) as exc:
        raise WorkbenchTermBridgeUnavailable("专业词库暂不可用") from exc
    if not isinstance(payload, dict):
        raise WorkbenchTermBridgeUnavailable("专业词库返回异常")
    return payload


def prepare_workbench_translation(
    source_path: str | Path,
    *,
    actor_sub: str,
    src_lang: str | None,
    tgt_lang: str | None,
    external_task_id: str | None = None,
) -> dict[str, Any]:
    path = Path(source_path)
    actor = (actor_sub or "").strip()
    if not actor:
        raise WorkbenchTermBridgeUnavailable("请登录后使用专业词库")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    source_text = _read_text(path) or path.stem
    run = _bridge_request(
        "POST",
        "/internal/workbench/v1/runs/start",
        {
            "actor_sub": actor,
            "source_sha256": digest,
            "source_text": source_text,
            "src_lang": _language(src_lang),
            "tgt_lang": _language(tgt_lang),
            "source_format": path.suffix.lstrip(".").casefold() or "unknown",
            "external_task_id": external_task_id,
        },
    )
    run["actor_sub"] = actor
    return run


def apply_pdf_term_policy(settings: Any, policy: dict | None, output_dir: str | Path) -> None:
    """Append only hard terms to BabelDOC's per-run CSV glossary list."""
    hard_terms = (policy or {}).get("hard_terms") or {}
    if not isinstance(hard_terms, dict) or not hard_terms:
        return
    directory = Path(output_dir) / ".qyunslation-termbase"
    directory.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(
        json.dumps(hard_terms, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()[:16]
    glossary_path = directory / f"terms-{digest}.csv"
    if not glossary_path.exists():
        with glossary_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["source", "target"])
            writer.writerows(sorted((str(source), str(target)) for source, target in hard_terms.items()))
    translation = getattr(settings, "translation", None)
    if translation is None:
        return
    current = str(getattr(translation, "glossaries", "") or "").strip()
    values = [value for value in current.split(",") if value]
    if str(glossary_path) not in values:
        values.append(str(glossary_path))
    translation.glossaries = ",".join(values)


def office_term_payload(policy: dict | None) -> dict[str, Any]:
    """Return additive payload fields already accepted by the Office sidecar."""
    return {
        "termbase_policy": policy or None,
        "termbase_version": (policy or {}).get("termbase_version"),
        "glossary_dict": dict((policy or {}).get("hard_terms") or {}),
    }


def _text_evidence(source: Path, target: Path) -> list[BilingualTermEvidence]:
    source_lines = _read_text(source).splitlines()
    target_lines = _read_text(target).splitlines()
    return [
        BilingualTermEvidence(
            source_text=source_line,
            target_text=target_lines[index] if index < len(target_lines) else "",
            block_id=f"line-{index + 1}",
        )
        for index, source_line in enumerate(source_lines[:1000])
        if source_line.strip()
    ]


def _pdf_evidence(source: Path, target: Path) -> list[BilingualTermEvidence]:
    import fitz

    rows: list[BilingualTermEvidence] = []
    with fitz.open(source) as source_pdf, fitz.open(target) as target_pdf:
        gate = ReferencePreserveGate()
        for index, source_page in enumerate(source_pdf):
            target_text = target_pdf[index].get_text("text").strip() if index < len(target_pdf) else ""
            blocks = sorted(
                source_page.get_text("blocks") or [],
                key=lambda block: (float(block[1]), float(block[0])),
            )
            for block_index, block in enumerate(blocks, start=1):
                if len(block) < 5 or not str(block[4]).strip():
                    continue
                source_text = str(block[4]).strip()
                role = "reference" if gate.should_preserve(source_text) else "body"
                rows.append(
                    BilingualTermEvidence(
                        source_text=source_text,
                        target_text=target_text,
                        role=role,
                        page_no=index + 1,
                        block_id=f"page-{index + 1}-block-{block_index}",
                        bbox={"x0": block[0], "y0": block[1], "x1": block[2], "y1": block[3]},
                    )
                )
    return rows


def _docx_evidence(source: Path, target: Path) -> list[BilingualTermEvidence]:
    from docx import Document

    source_doc, target_doc = Document(source), Document(target)
    rows: list[BilingualTermEvidence] = []
    gate = ReferencePreserveGate()
    for index, paragraph in enumerate(source_doc.paragraphs):
        if paragraph.text.strip():
            target_text = target_doc.paragraphs[index].text if index < len(target_doc.paragraphs) else ""
            role = "reference" if gate.should_preserve(paragraph.text) else "body"
            rows.append(BilingualTermEvidence(source_text=paragraph.text, target_text=target_text, role=role, block_id=f"paragraph-{index + 1}"))
    for table_index, table in enumerate(source_doc.tables):
        target_table = target_doc.tables[table_index] if table_index < len(target_doc.tables) else None
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                if not cell.text.strip():
                    continue
                target_text = ""
                if target_table and row_index < len(target_table.rows) and cell_index < len(target_table.rows[row_index].cells):
                    target_text = target_table.rows[row_index].cells[cell_index].text
                source_term, target_term = _cell_term_pair(cell.text, target_text)
                rows.append(BilingualTermEvidence(source_text=cell.text, target_text=target_text, role="table", block_id=f"table-{table_index + 1}-r{row_index + 1}-c{cell_index + 1}", object_id=f"table-{table_index + 1}", source_term=source_term, target_term=target_term))
    rows.extend(_docx_footnote_evidence(source, target))
    return rows[:1000]


def _pptx_evidence(source: Path, target: Path) -> list[BilingualTermEvidence]:
    from pptx import Presentation

    source_deck, target_deck = Presentation(source), Presentation(target)
    rows: list[BilingualTermEvidence] = []
    for slide_index, slide in enumerate(source_deck.slides):
        target_slide = target_deck.slides[slide_index] if slide_index < len(target_deck.slides) else None
        for shape_index, shape in enumerate(slide.shapes):
            target_shape = target_slide.shapes[shape_index] if target_slide and shape_index < len(target_slide.shapes) else None
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                rows.append(BilingualTermEvidence(source_text=shape.text, target_text=getattr(target_shape, "text", "") if target_shape else "", page_no=slide_index + 1, block_id=f"slide-{slide_index + 1}-shape-{shape_index + 1}", object_id=f"slide-{slide_index + 1}-shape-{shape_index + 1}"))
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows):
                    for cell_index, cell in enumerate(row.cells):
                        if cell.text.strip():
                            target_text = ""
                            if target_shape and getattr(target_shape, "has_table", False):
                                target_text = target_shape.table.rows[row_index].cells[cell_index].text
                            source_term, target_term = _cell_term_pair(cell.text, target_text)
                            rows.append(BilingualTermEvidence(source_text=cell.text, target_text=target_text, role="table", page_no=slide_index + 1, block_id=f"slide-{slide_index + 1}-table-{shape_index + 1}-r{row_index + 1}-c{cell_index + 1}", object_id=f"slide-{slide_index + 1}-table-{shape_index + 1}", source_term=source_term, target_term=target_term))
    return rows[:1000]


def _pptx_embedded_image_evidence(source: Path, target: Path) -> tuple[list[BilingualTermEvidence], str | None]:
    try:
        from pptx import Presentation

        source_deck, target_deck = Presentation(source), Presentation(target)
    except Exception:
        return [], "pptx_embedded_image_adapter_failed"
    rows: list[BilingualTermEvidence] = []
    reasons: list[str] = []
    for slide_index, source_slide in enumerate(source_deck.slides, start=1):
        target_slide = target_deck.slides[slide_index - 1] if slide_index <= len(target_deck.slides) else None
        for shape_index, source_shape in enumerate(source_slide.shapes, start=1):
            try:
                source_blob = source_shape.image.blob
            except Exception:
                continue
            target_shape = target_slide.shapes[shape_index - 1] if target_slide and shape_index <= len(target_slide.shapes) else None
            try:
                target_blob = target_shape.image.blob if target_shape else None
            except Exception:
                target_blob = None
            if not target_blob:
                reasons.append("pptx_embedded_image_count_or_shape_mismatch")
                continue
            image_rows, reason = _image_blob_evidence(
                source_blob,
                target_blob,
                object_id=f"slide-{slide_index}-image-{shape_index}",
            )
            rows.extend(image_rows)
            if reason:
                reasons.append(reason)
    return rows, ";".join(sorted(set(reasons))) or None


def _pptx_evidence_with_status(source: Path, target: Path) -> tuple[list[BilingualTermEvidence], str | None]:
    rows = _pptx_evidence(source, target)
    image_rows, reason = _pptx_embedded_image_evidence(source, target)
    return rows + image_rows, reason


def _docx_footnotes(path: Path) -> dict[str, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("word/footnotes.xml"))
    except (KeyError, OSError, zipfile.BadZipFile, ElementTree.ParseError):
        return {}
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    rows: dict[str, str] = {}
    for node in root.findall(f"{namespace}footnote"):
        identifier = node.attrib.get(f"{namespace}id")
        if identifier is None:
            continue
        try:
            if int(identifier) < 0:
                continue
        except ValueError:
            continue
        text = "".join(node.itertext()).strip()
        if text:
            rows[identifier] = text
    return rows


def _docx_footnote_evidence(source: Path, target: Path) -> list[BilingualTermEvidence]:
    source_notes, target_notes = _docx_footnotes(source), _docx_footnotes(target)
    rows: list[BilingualTermEvidence] = []
    gate = ReferencePreserveGate()
    for identifier, source_text in source_notes.items():
        target_text = target_notes.get(identifier, "")
        source_term, target_term = _cell_term_pair(source_text, target_text)
        rows.append(
            BilingualTermEvidence(
                source_text=source_text,
                target_text=target_text,
                role="reference" if gate.should_preserve(source_text) else "footnote",
                block_id=f"footnote-{identifier}",
                object_id="footnotes",
                source_term=source_term,
                target_term=target_term,
            )
        )
    return rows


def _docx_embedded_image_evidence(source: Path, target: Path) -> tuple[list[BilingualTermEvidence], str | None]:
    """Pair DOCX DrawingML occurrences by document occurrence order.

    DOCX writeback preserves occurrence order even when a source blob is reused
    by more than one drawing.  A count mismatch is visible as degradation.
    """
    try:
        from docx import Document
        from qyunslation.extensions.docx_image_overlay import enumerate_drawing_occurrences

        source_images = enumerate_drawing_occurrences(Document(source))
        target_images = enumerate_drawing_occurrences(Document(target))
    except Exception:
        return [], "docx_embedded_image_adapter_failed"
    rows: list[BilingualTermEvidence] = []
    reasons: list[str] = []
    for index, source_image in enumerate(source_images):
        if index >= len(target_images):
            reasons.append("docx_embedded_image_count_mismatch")
            break
        try:
            image_rows, reason = _image_blob_evidence(
                source_image.image_part.blob,
                target_images[index].image_part.blob,
                object_id=f"docx-image-{index + 1}",
            )
            rows.extend(image_rows)
            if reason:
                reasons.append(reason)
        except Exception:
            reasons.append("docx_embedded_image_adapter_failed")
    if len(target_images) != len(source_images):
        reasons.append("docx_embedded_image_count_mismatch")
    return rows, ";".join(sorted(set(reasons))) or None


def _docx_evidence_with_status(source: Path, target: Path) -> tuple[list[BilingualTermEvidence], str | None]:
    rows = _docx_evidence(source, target)
    image_rows, reason = _docx_embedded_image_evidence(source, target)
    return rows + image_rows, reason


def _image_evidence(source: Path, target: Path, *, object_id: str = "image-1") -> tuple[list[BilingualTermEvidence], str | None]:
    """Pair source and translated OCR boxes by stable overlay geometry.

    The image-overlay workflow redraws in the original box.  If that invariant
    is not observable we report a degradation instead of inventing a bilingual
    pair from reading order alone.
    """
    from qyunslation.extensions.image_translate import ocr_image

    try:
        source_boxes = ocr_image(source)
        target_boxes = ocr_image(target)
        from PIL import Image

        with Image.open(source) as image:
            diagonal = math.hypot(*image.size) or 1.0
    except Exception:
        return [], "image_ocr_failed"
    if not source_boxes:
        return [], "image_ocr_no_source_text"
    if not target_boxes:
        return [], "image_ocr_no_translated_text"

    available = set(range(len(target_boxes)))
    rows: list[BilingualTermEvidence] = []
    for index, source_box in enumerate(source_boxes, start=1):
        sx0, sy0, sx1, sy1, source_text, _score = source_box
        source_text = str(source_text or "").strip()
        if not source_text:
            continue
        source_center = ((sx0 + sx1) / 2, (sy0 + sy1) / 2)
        best_index = min(
            available,
            key=lambda candidate: math.dist(
                source_center,
                ((target_boxes[candidate][0] + target_boxes[candidate][2]) / 2, (target_boxes[candidate][1] + target_boxes[candidate][3]) / 2),
            ),
            default=None,
        )
        if best_index is None:
            continue
        target_box = target_boxes[best_index]
        distance = math.dist(
            source_center,
            ((target_box[0] + target_box[2]) / 2, (target_box[1] + target_box[3]) / 2),
        )
        if distance > diagonal * 0.12:
            continue
        available.remove(best_index)
        target_text = str(target_box[4] or "").strip()
        if not target_text:
            continue
        rows.append(
            BilingualTermEvidence(
                source_text=source_text,
                target_text=target_text,
                role="ocr",
                block_id=f"ocr-{index}",
                object_id=object_id,
                bbox={"x0": sx0, "y0": sy0, "x1": sx1, "y1": sy1},
                source_term=source_text[:512],
                target_term=target_text[:512],
            )
        )
    return (rows, None) if rows and len(rows) == len(source_boxes) else (rows, "image_ocr_alignment_incomplete")


def _image_blob_evidence(source_blob: bytes, target_blob: bytes, *, object_id: str) -> tuple[list[BilingualTermEvidence], str | None]:
    """Run the same evidence adapter for Office embedded raster images."""
    with tempfile.TemporaryDirectory(prefix="qyunslation-term-evidence-") as directory:
        root = Path(directory)
        source = root / "source.png"
        target = root / "target.png"
        source.write_bytes(source_blob)
        target.write_bytes(target_blob)
        return _image_evidence(source, target, object_id=object_id)


def build_bilingual_evidence(source_path: str | Path, target_path: str | Path) -> tuple[list[BilingualTermEvidence], str | None]:
    """Build format-aware evidence; never pretend unsupported binary paths are empty."""
    source, target = Path(source_path), Path(target_path)
    try:
        suffix = source.suffix.casefold()
        if suffix == ".pdf" and target.suffix.casefold() == ".pdf":
            rows = _pdf_evidence(source, target)
        elif suffix == ".docx" and target.suffix.casefold() == ".docx":
            return _docx_evidence_with_status(source, target)
        elif suffix in {".ppt", ".pptx"} and target.suffix.casefold() == ".pptx":
            return _pptx_evidence_with_status(source, target)
        elif suffix in _IMAGE_SUFFIXES and target.suffix.casefold() in _IMAGE_SUFFIXES:
            return _image_evidence(source, target)
        elif suffix in _TEXT_SUFFIXES and target.suffix.casefold() in _TEXT_SUFFIXES:
            rows = _text_evidence(source, target)
        else:
            return [], "unsupported_or_unaligned_output"
    except Exception:
        return [], "evidence_adapter_failed"
    return (rows, None) if rows else ([], "no_aligned_evidence")


def complete_workbench_translation(run: dict[str, Any], source_path: str | Path, target_path: str | Path | None) -> dict[str, Any]:
    if not target_path:
        return _bridge_request(
            "POST",
            f"/internal/workbench/v1/runs/{run['run_id']}/complete",
            {"actor_sub": run["actor_sub"], "evidence": [], "degradation_reason": "translated_output_missing"},
        )
    evidence, reason = build_bilingual_evidence(source_path, target_path)
    return _bridge_request(
        "POST",
        f"/internal/workbench/v1/runs/{run['run_id']}/complete",
        {
            "actor_sub": run["actor_sub"],
            "evidence": [
                {
                    "source_text": item.source_text,
                    "target_text": item.target_text,
                    "role": item.role,
                    "page_no": item.page_no,
                    "block_id": item.block_id,
                    "object_id": item.object_id,
                    "char_start": item.char_start,
                    "char_end": item.char_end,
                    "bbox": item.bbox,
                    "source_term": item.source_term,
                    "target_term": item.target_term,
                    "term_type": item.term_type,
                }
                for item in evidence
            ],
            "degradation_reason": reason,
        },
    )


def restore_latest_workbench_run(actor_sub: str) -> dict[str, Any] | None:
    """Reload the actor's latest reviewable run after Gradio state is wiped."""
    actor = (actor_sub or "").strip()
    if not actor:
        return None
    try:
        payload = _bridge_request(
            "GET",
            f"/internal/workbench/v1/runs/latest?actor_sub={actor}",
            {},
        )
    except WorkbenchTermBridgeUnavailable as exc:
        if exc.status_code == 404:
            return None
        raise
    run_id = payload.get("run_id")
    if not run_id:
        return None
    return {
        "run_id": run_id,
        "job_id": payload.get("job_id"),
        "actor_sub": actor,
        "policy": payload.get("policy") or {},
        "restored": True,
    }


def search_workbench_concepts(run: dict[str, Any], source_term: str) -> list[dict[str, Any]]:
    payload = _bridge_request(
        "POST",
        "/internal/workbench/v1/concepts/search",
        {"actor_sub": run["actor_sub"], "query": (source_term or "").strip() or "term"},
    )
    matches = payload.get("matches")
    return matches if isinstance(matches, list) else []


def get_workbench_term_review(run: dict[str, Any]) -> dict[str, Any]:
    return _bridge_request(
        "GET",
        f"/internal/workbench/v1/runs/{run['run_id']}/term-review?actor_sub={run['actor_sub']}",
        {},
    )


def decide_workbench_term(run: dict[str, Any], candidate_id: str, decision: dict[str, Any]) -> dict[str, Any]:
    body = dict(decision)
    body["actor_sub"] = run["actor_sub"]
    return _bridge_request(
        "POST",
        f"/internal/workbench/v1/runs/{run['run_id']}/terms/{candidate_id}/decision",
        body,
    )


def batch_decide_workbench_terms(run: dict[str, Any], decisions: list[dict[str, Any]]) -> dict[str, Any]:
    """Approve only server-eligible low-risk exact or alias candidates."""
    return _bridge_request(
        "POST",
        f"/internal/workbench/v1/runs/{run['run_id']}/terms/batch-decision",
        {"actor_sub": run["actor_sub"], "decisions": decisions},
    )
