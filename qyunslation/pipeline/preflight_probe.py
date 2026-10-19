# SPDX-License-Identifier: MPL-2.0
"""PLAN-078f f2：预检时对原文件做一次有时间上限的结构探测。

预检卡要告诉人“这份文件翻译起来有什么限制”：是否加密、有多少扫描页、有没有复杂表格、
语言和所选源语言是否对得上。这里只读文件、不调用任何外部服务；任何一步失败都降级为
一条提示（不影响继续），只有“确实无法翻译”的情况（加密、损坏、空文件）才是阻断项。

每条提示都写清三件事：发现了什么（message）、对结果的影响（impact）、建议怎么做（action）。
"""
from __future__ import annotations

import logging
import os
import re
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

PROBE_BUDGET_SECONDS = 4.0
PDF_SAMPLE_PAGES = 40
TABLE_SAMPLE_PAGES = 8
SCANNED_PAGE_CHARS = 25  # a page with fewer characters than this and an image is a scan
MIN_LETTERS_FOR_LANGUAGE = 200
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # encrypted Office files are OLE containers
LOW_RESOLUTION_PX = 800

_CJK = re.compile(r"[一-鿿]")
_LATIN = re.compile(r"[A-Za-z]")


@dataclass
class ProbeResult:
    manifest_summary: dict[str, Any] = field(default_factory=dict)
    issues: list[dict[str, str]] = field(default_factory=list)

    @property
    def blocked(self) -> bool:
        return any(item["severity"] == "blocker" for item in self.issues)


def _issue(severity: str, code: str, message: str, impact: str, action: str) -> dict[str, str]:
    return {"severity": severity, "code": code, "message": message, "impact": impact, "action": action}


def _language_issue(text: str, source_language: str | None) -> dict[str, str] | None:
    cjk = len(_CJK.findall(text))
    latin = len(_LATIN.findall(text))
    if cjk + latin < MIN_LETTERS_FOR_LANGUAGE:
        return None  # too little text to say anything
    cjk_share = cjk / (cjk + latin)
    if source_language == "English" and cjk_share > 0.5:
        return _issue(
            "warning", "LANGUAGE_MISMATCH",
            f"文件内容看起来主要是中文（约 {round(cjk_share * 100)}%），但源语言选的是 English。",
            "按英文翻译中文原文，译文会出现大量原样保留或错译。",
            "确认语言方向；如果原文确实是中文，请改为“简体中文 → English”。",
        )
    if source_language == "简体中文" and cjk_share < 0.1:
        return _issue(
            "warning", "LANGUAGE_MISMATCH",
            f"文件内容看起来主要是外文（中文只占约 {round(cjk_share * 100)}%），但源语言选的是简体中文。",
            "按中文翻译外文原文，译文会出现大量原样保留或错译。",
            "确认语言方向；如果原文确实是英文，请改为“English → 简体中文”。",
        )
    return None


def _encrypted(kind: str) -> dict[str, str]:
    return _issue(
        "blocker", f"{kind}_ENCRYPTED",
        "文件有打开密码，系统无法读取内容。",
        "无法翻译。",
        "请在原软件里去掉密码保护后重新上传。",
    )


def _unreadable(kind: str) -> dict[str, str]:
    return _issue(
        "blocker", f"{kind}_UNREADABLE",
        "文件已损坏或不是有效的文档，系统无法打开。",
        "无法翻译。",
        "请用原软件重新另存一份后再上传。",
    )


def _is_ole(path: Path) -> bool:
    try:
        with path.open("rb") as handle:
            return handle.read(8) == OLE_MAGIC
    except OSError:
        return False


def _probe_pdf(path: Path, source_language: str | None, deadline: float) -> ProbeResult:
    import pymupdf

    result = ProbeResult()
    try:
        doc = pymupdf.open(str(path))
    except Exception:
        result.issues.append(_unreadable("PDF"))
        return result
    try:
        if doc.needs_pass:
            result.issues.append(_encrypted("PDF"))
            return result
        pages = doc.page_count
        if pages <= 0:
            result.issues.append(_issue("blocker", "PDF_EMPTY", "文件没有任何页面。", "无法翻译。", "请检查原文件后重新上传。"))
            return result
        step = max(1, pages // PDF_SAMPLE_PAGES)
        indices = list(range(0, pages, step))[:PDF_SAMPLE_PAGES]
        sampled = scanned = images = 0
        texts: list[str] = []
        tables = max_columns = 0
        for position, index in enumerate(indices):
            if time.monotonic() > deadline:
                break
            page = doc[index]
            text = page.get_text("text") or ""
            page_images = len(page.get_images(full=False))
            sampled += 1
            images += page_images
            if len(text.strip()) < SCANNED_PAGE_CHARS and page_images >= 1:
                scanned += 1
            else:
                texts.append(text)
            if position < TABLE_SAMPLE_PAGES and time.monotonic() <= deadline:
                try:
                    for table in page.find_tables().tables:
                        tables += 1
                        max_columns = max(max_columns, table.col_count)
                except Exception:
                    pass  # table detection is best-effort
        ratio = scanned / sampled if sampled else 0
        estimated_scanned = round(ratio * pages)
        result.manifest_summary = {
            "pages": pages,
            "sampled_pages": sampled,
            "scanned_pages": estimated_scanned,
            "scan_detected": ratio >= 0.5,
            "tables": tables or None,
            "images": images or None,
            "objects": None,
        }
        if ratio >= 0.5:
            result.issues.append(_issue(
                "warning", "SCANNED_PDF",
                f"文件主要由扫描页构成（约 {estimated_scanned}/{pages} 页没有文字层）。",
                "需要先做 OCR，耗时明显更长；识别错误会带进译文，版式还原也更容易出偏差。",
                "尽量提供带文字层的原件；继续的话系统会自动补充 OCR 文字层，请在质量检查里重点核对数字和术语。",
            ))
        elif scanned:
            result.issues.append(_issue(
                "warning", "SCANNED_PAGES",
                f"约有 {estimated_scanned} 页是扫描页（没有文字层）。",
                "这些页的文字要靠 OCR 识别，准确度低于有文字层的页面。",
                "翻译后重点核对这几页；如有可能请换用带文字层的原件。",
            ))
        if max_columns >= 8:
            result.issues.append(_issue(
                "warning", "COMPLEX_TABLES",
                f"检测到较宽的表格（最多 {max_columns} 列，共 {tables} 个表格）。",
                "宽表格在译文中容易换行错位或溢出。",
                "翻译后核对表格版式；必要时拆分表格后再翻译。",
            ))
        elif tables:
            result.issues.append(_issue(
                "info", "TABLES",
                f"检测到 {tables} 个表格。", "表格内的数字和单位会进入质量检查。", "无需处理。",
            ))
        if images >= 10:
            result.issues.append(_issue(
                "info", "MANY_IMAGES",
                f"抽查的页面中含有 {images} 张图片。",
                "图片里的文字不一定会被翻译，也不一定能被质量检查覆盖。",
                "翻译后检查图片中的文字是否需要处理。",
            ))
        language = _language_issue("\n".join(texts), source_language)
        if language:
            result.issues.append(language)
        return result
    finally:
        doc.close()


def _probe_docx(path: Path, source_language: str | None, deadline: float) -> ProbeResult:
    result = ProbeResult()
    if not zipfile.is_zipfile(path):
        result.issues.append(_encrypted("DOCX") if _is_ole(path) else _unreadable("DOCX"))
        return result
    import docx

    try:
        document = docx.Document(str(path))
    except Exception:
        result.issues.append(_unreadable("DOCX"))
        return result
    tables = list(document.tables)
    nested = any(cell.tables for table in tables for row in table.rows for cell in row.cells) if len(tables) <= 40 else False
    widest = max((len(table.columns) for table in tables), default=0)
    images = len(document.inline_shapes)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)[:200_000]
    result.manifest_summary = {
        "pages": None,
        "paragraphs": len(document.paragraphs),
        "tables": len(tables) or None,
        "images": images or None,
        "objects": None,
        "scan_detected": False,
    }
    if nested or widest >= 8:
        result.issues.append(_issue(
            "warning", "COMPLEX_TABLES",
            "文件含有嵌套表格或很宽的表格" + (f"（最多 {widest} 列）" if widest >= 8 else "") + "。",
            "这类表格在译文中容易错位，需要人工核对版式。",
            "翻译后核对表格；必要时简化表格结构再翻译。",
        ))
    elif tables:
        result.issues.append(_issue("info", "TABLES", f"检测到 {len(tables)} 个表格。", "表格内的数字和单位会进入质量检查。", "无需处理。"))
    if images >= 10:
        result.issues.append(_issue(
            "info", "MANY_IMAGES", f"文件含有 {images} 张内嵌图片。",
            "图片里的文字不一定会被翻译，也不一定能被质量检查覆盖。", "翻译后检查图片中的文字是否需要处理。",
        ))
    language = _language_issue(text, source_language)
    if language:
        result.issues.append(language)
    return result


def _probe_pptx(path: Path, source_language: str | None, deadline: float) -> ProbeResult:
    result = ProbeResult()
    if not zipfile.is_zipfile(path):
        result.issues.append(_encrypted("PPTX") if _is_ole(path) else _unreadable("PPTX"))
        return result
    from pptx import Presentation

    try:
        deck = Presentation(str(path))
    except Exception:
        result.issues.append(_unreadable("PPTX"))
        return result
    slides = len(deck.slides)
    pictures = tables = 0
    texts: list[str] = []
    for slide in deck.slides:
        if time.monotonic() > deadline:
            break
        for shape in slide.shapes:
            if shape.shape_type == 13:  # picture
                pictures += 1
            if getattr(shape, "has_table", False) and shape.has_table:
                tables += 1
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                texts.append(shape.text_frame.text)
    result.manifest_summary = {
        "pages": slides,
        "tables": tables or None,
        "images": pictures or None,
        "objects": None,
        "scan_detected": False,
    }
    if slides == 0:
        result.issues.append(_issue("blocker", "PPTX_EMPTY", "演示文稿里没有幻灯片。", "无法翻译。", "请检查原文件后重新上传。"))
        return result
    if pictures >= 10:
        result.issues.append(_issue(
            "info", "MANY_IMAGES", f"文件含有 {pictures} 张图片。",
            "图片里的文字不一定会被翻译，也不一定能被质量检查覆盖。", "翻译后检查图片中的文字是否需要处理。",
        ))
    language = _language_issue("\n".join(texts)[:200_000], source_language)
    if language:
        result.issues.append(language)
    return result


def _probe_text(path: Path, source_language: str | None) -> ProbeResult:
    result = ProbeResult()
    raw = path.read_bytes()[:500_000]
    text = raw.decode("utf-8", errors="replace")
    result.manifest_summary = {"pages": None, "characters": len(text), "objects": None, "scan_detected": False}
    if not text.strip():
        result.issues.append(_issue("blocker", "TEXT_EMPTY", "文件里没有可翻译的文字。", "无法翻译。", "请检查原文件后重新上传。"))
        return result
    language = _language_issue(text, source_language)
    if language:
        result.issues.append(language)
    return result


def _probe_image(path: Path) -> ProbeResult:
    from PIL import Image

    result = ProbeResult()
    try:
        with Image.open(path) as image:
            width, height = image.size
    except Exception:
        result.issues.append(_unreadable("IMAGE"))
        return result
    result.manifest_summary = {"pages": 1, "width": width, "height": height, "objects": None, "scan_detected": True}
    if max(width, height) < LOW_RESOLUTION_PX:
        result.issues.append(_issue(
            "warning", "LOW_RESOLUTION",
            f"图片分辨率偏低（{width}×{height}）。",
            "OCR 对小字和数字容易识别错，错误会带进译文。",
            "请提供更清晰的原图（长边建议不小于 1200 像素）。",
        ))
    return result


def probe_document(path: Path, suffix: str, source_language: str | None) -> ProbeResult:
    """Inspect ``path`` within a small time budget. Never raises.

    ``QYUNSLATION_PREFLIGHT_PROBE=0`` switches the probe off (an operator escape hatch if a parser
    misbehaves on real files; the test suite also uses it, because most fixtures upload placeholder bytes).
    """
    if os.environ.get("QYUNSLATION_PREFLIGHT_PROBE", "1").strip().casefold() in {"0", "false", "off", "no"}:
        return ProbeResult()
    deadline = time.monotonic() + PROBE_BUDGET_SECONDS
    kind = suffix.casefold().lstrip(".")
    try:
        if kind == "pdf":
            return _probe_pdf(path, source_language, deadline)
        if kind == "docx":
            return _probe_docx(path, source_language, deadline)
        if kind == "pptx":
            return _probe_pptx(path, source_language, deadline)
        if kind in {"txt", "md"}:
            return _probe_text(path, source_language)
        if kind in {"png", "jpg", "jpeg"}:
            return _probe_image(path)
    except Exception:
        logger.warning("preflight probe failed for %s", path.name, exc_info=True)
        return ProbeResult(
            issues=[_issue(
                "info", "PROBE_FAILED",
                "没能分析文件结构。", "预检卡里看不到页数、扫描页等信息，不影响继续翻译。", "无需处理。",
            )]
        )
    return ProbeResult()
