# SPDX-License-Identifier: MPL-2.0
"""PLAN-030g：PPTX 结构扫描与图片化打包。"""
from __future__ import annotations

import hashlib
import shutil
import subprocess
import tempfile
from collections.abc import Callable
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from .ingest import InputPreparationError, prepare_document
from .models import (
    BlockRole,
    CURRENT_SCHEMA_VERSION,
    AssetRef,
    AssetRole,
    BoundingBox,
    ContentProfile,
    DetectorEvidence,
    DocumentInfo,
    DocumentStructureManifest,
    ImageObject,
    ObjectType,
    OutputEditability,
    ProcessingMode,
    ProducerInfo,
    Representation,
    SourceFormat,
    SourceRef,
    SourceRefKind,
    TableObject,
    TextBoxObject,
    TranslatableBlock,
    build_manifest_id,
    build_object_id,
)
from .profiles import resolve_profile
from .table_cell_policy import classify_cell_policy

EMU_PER_PT = 12_700.0
PPTX_STRUCTURE_SCANNER_NAME = "qyunslation-plan-030g"
PPTX_STRUCTURE_SCANNER_VERSION = "1.0.0"
SlideRenderer = Callable[[bytes], list[bytes]]


def _emu_bbox(shape) -> BoundingBox:
    x0 = max(0.0, float(shape.left) / EMU_PER_PT)
    y0 = max(0.0, float(shape.top) / EMU_PER_PT)
    width = max(float(shape.width) / EMU_PER_PT, 0.01)
    height = max(float(shape.height) / EMU_PER_PT, 0.01)
    return BoundingBox(x0=x0, y0=y0, x1=x0 + width, y1=y0 + height)


def _canvas_bbox(canvas) -> BoundingBox:
    return BoundingBox(x0=0.0, y0=0.0, x1=float(canvas.width), y1=float(canvas.height))


def _walk_shapes(shape, pictures: list, tables: list, textboxes: list) -> None:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            _walk_shapes(child, pictures, tables, textboxes)
        return
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        pictures.append(shape)
        return
    if shape.has_table:
        tables.append(shape)
        return
    if getattr(shape, "has_text_frame", False) and shape.text_frame:
        if (shape.text_frame.text or "").strip():
            textboxes.append(shape)


def render_pptx_slides(content: bytes, *, renderer: SlideRenderer | None = None) -> list[bytes]:
    """Render each slide to PNG. Missing renderer is fail-closed."""
    if renderer is not None:
        pages = renderer(content)
        if not pages:
            raise InputPreparationError(
                "SLIDE_RENDER_EMPTY",
                "幻灯片渲染器未产出任何页图",
                http_status=503,
            )
        return list(pages)

    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise InputPreparationError(
            "RUNTIME_CAPABILITY_UNAVAILABLE",
            "当前运行环境不能处理 PPTX（缺少：SLIDE_RENDERER）",
            http_status=503,
        )

    with tempfile.TemporaryDirectory(prefix="qy_pptx_render_") as temp_name:
        temp_dir = Path(temp_name)
        source = temp_dir / "deck.pptx"
        source.write_bytes(content)
        command = [
            executable,
            "--headless",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            str(temp_dir),
            str(source),
        ]
        try:
            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise InputPreparationError(
                "SLIDE_RENDER_FAILED",
                "幻灯片渲染失败",
                http_status=503,
            ) from exc
        pdf_path = temp_dir / "deck.pdf"
        if completed.returncode != 0 or not pdf_path.is_file():
            raise InputPreparationError(
                "SLIDE_RENDER_FAILED",
                "幻灯片渲染失败",
                http_status=503,
            )
        import pymupdf

        document = pymupdf.open(pdf_path)
        try:
            pages = [page.get_pixmap(dpi=144).tobytes("png") for page in document]
        finally:
            document.close()
        if not pages:
            raise InputPreparationError(
                "SLIDE_RENDER_EMPTY",
                "幻灯片渲染器未产出任何页图",
                http_status=503,
            )
        return pages


def pack_image_pptx(
    pages: list[bytes],
    *,
    width_pt: float,
    height_pt: float,
) -> bytes:
    """Pack overlayed page images onto a blank-slide PPTX backboard."""
    if not pages:
        raise ValueError("pack_image_pptx requires at least one page image")
    presentation = Presentation()
    presentation.slide_width = Emu(int(width_pt * EMU_PER_PT))
    presentation.slide_height = Emu(int(height_pt * EMU_PER_PT))
    blank = presentation.slide_layouts[6]
    for page in pages:
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(
            BytesIO(page),
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


class PptxStructureScanner:
    def __init__(self, *, slide_renderer: SlideRenderer | None = None) -> None:
        self.slide_renderer = slide_renderer

    def scan(
        self,
        source: Path | str | bytes,
        *,
        source_name: str = "deck.pptx",
        content_profile: ContentProfile | None = None,
        processing_mode: ProcessingMode | None = None,
    ) -> DocumentStructureManifest:
        if isinstance(source, (Path, str)):
            path = Path(source)
            content = path.read_bytes()
            source_name = path.name
        else:
            content = bytes(source)

        prepared = prepare_document(source_name, content)
        if prepared.normalized_format is not SourceFormat.PPTX:
            raise ValueError("PptxStructureScanner expects PPT or PPTX")

        mode = processing_mode or ProcessingMode.NATIVE
        selected = (
            ProcessingMode.RENDERED if mode is ProcessingMode.RENDERED else ProcessingMode.NATIVE
        )
        canvases = {canvas.source_index: canvas for canvas in prepared.canvases}
        objects: list = []
        derived_assets = list(prepared.derived_assets)
        pictures = 0
        physical_image_keys: set[str] = set()

        if selected is ProcessingMode.RENDERED:
            pages = render_pptx_slides(prepared.content, renderer=self.slide_renderer)
            if len(pages) != len(prepared.canvases):
                raise InputPreparationError(
                    "SLIDE_RENDER_COUNT_MISMATCH",
                    "渲染页数与幻灯片画布不一致",
                    http_status=503,
                )
            for canvas, page in zip(prepared.canvases, pages, strict=True):
                digest = hashlib.sha256(page).hexdigest()
                derived_assets.append(
                    AssetRef(
                        asset_id=f"asset:rendered:{digest}",
                        role=AssetRole.RENDERED,
                        sha256=digest,
                        media_type="image/png",
                        locator=f"rendered://{canvas.canvas_id}",
                    )
                )
                key = f"image:{canvas.canvas_id}"
                object_id = build_object_id(
                    prepared.source_sha256,
                    canvas.canvas_id,
                    ObjectType.IMAGE,
                    key,
                    [{"kind": SourceRefKind.PPTX_SLIDE.value, "ref": canvas.canvas_id}],
                )
                objects.append(
                    ImageObject(
                        type=ObjectType.IMAGE,
                        object_id=object_id,
                        canvas_id=canvas.canvas_id,
                        bbox=_canvas_bbox(canvas),
                        representation=Representation.BITMAP,
                        source_refs=[
                            SourceRef(kind=SourceRefKind.PPTX_SLIDE, ref=canvas.canvas_id)
                        ],
                        planned_action="ocr_overlay",
                        semantic_id=key,
                        occurrence_key=key,
                    )
                )
                canvas.reading_order = [object_id]
            pictures = len(pages)
            physical_image_keys = {f"rendered:slide:{index}" for index in range(1, len(pages) + 1)}
        else:
            presentation = Presentation(BytesIO(prepared.content))
            for index, slide in enumerate(presentation.slides, start=1):
                canvas = canvases.get(index)
                if canvas is None:
                    continue
                found_pictures: list = []
                found_tables: list = []
                found_textboxes: list = []
                for shape in slide.shapes:
                    _walk_shapes(shape, found_pictures, found_tables, found_textboxes)
                for shape in found_tables:
                    table = shape.table
                    key = f"table:slide:{index}:{shape.shape_id}"
                    object_id = build_object_id(
                        prepared.source_sha256,
                        canvas.canvas_id,
                        ObjectType.TABLE,
                        key,
                        [{"kind": SourceRefKind.PPTX_SHAPE.value, "ref": key}],
                    )
                    blocks = []
                    for row_i, row in enumerate(table.rows):
                        for col_i, cell in enumerate(row.cells):
                            text = (cell.text or "").strip()
                            if text:
                                blocks.append(
                                    TranslatableBlock(
                                        block_id=f"{key}:r{row_i}c{col_i}",
                                        source_text=text,
                                        role=BlockRole.TABLE_CELL,
                                        translation_policy=classify_cell_policy(text),
                                        row_index=row_i,
                                        column_index=col_i,
                                    )
                                )
                    objects.append(
                        TableObject(
                            type=ObjectType.TABLE,
                            object_id=object_id,
                            canvas_id=canvas.canvas_id,
                            bbox=_emu_bbox(shape),
                            representation=Representation.NATIVE_TEXT,
                            source_refs=[
                                SourceRef(kind=SourceRefKind.PPTX_SHAPE, ref=key)
                            ],
                            translatable_blocks=blocks,
                            planned_action="native_text",
                            row_count=len(table.rows),
                            column_count=len(table.columns),
                            semantic_id=key,
                        )
                    )
                for shape in found_textboxes:
                    text = (shape.text_frame.text or "").strip()
                    key = f"textbox:slide:{index}:{shape.shape_id}"
                    object_id = build_object_id(
                        prepared.source_sha256,
                        canvas.canvas_id,
                        ObjectType.TEXT_BOX,
                        key,
                        [{"kind": SourceRefKind.PPTX_SHAPE.value, "ref": key}],
                    )
                    objects.append(
                        TextBoxObject(
                            type=ObjectType.TEXT_BOX,
                            object_id=object_id,
                            canvas_id=canvas.canvas_id,
                            bbox=_emu_bbox(shape),
                            representation=Representation.NATIVE_TEXT,
                            source_refs=[
                                SourceRef(kind=SourceRefKind.PPTX_SHAPE, ref=key)
                            ],
                            translatable_blocks=[
                                TranslatableBlock(block_id=f"{key}:0", source_text=text)
                            ],
                            planned_action="native_text",
                            semantic_id=key,
                        )
                    )
                for pic_i, pic in enumerate(found_pictures, start=1):
                    pictures += 1
                    try:
                        physical_image_keys.add(
                            f"blob:{hashlib.sha256(pic.image.blob).hexdigest()}"
                        )
                    except Exception:
                        physical_image_keys.add(f"shape:{index}:{pic.shape_id}")
                    key = f"image:slide:{index}:{pic.shape_id}"
                    object_id = build_object_id(
                        prepared.source_sha256,
                        canvas.canvas_id,
                        ObjectType.IMAGE,
                        key,
                        [{"kind": SourceRefKind.PPTX_SHAPE.value, "ref": key}],
                    )
                    objects.append(
                        ImageObject(
                            type=ObjectType.IMAGE,
                            object_id=object_id,
                            canvas_id=canvas.canvas_id,
                            bbox=_emu_bbox(pic),
                            representation=Representation.BITMAP,
                            source_refs=[
                                SourceRef(
                                    kind=SourceRefKind.PPTX_SHAPE,
                                    ref=key,
                                    occurrence_index=pic_i,
                                )
                            ],
                            detector_evidence=[
                                DetectorEvidence(
                                    detector="python-pptx",
                                    label="picture",
                                    bbox=_emu_bbox(pic),
                                    details={"slide": index, "shape_id": pic.shape_id},
                                )
                            ],
                            planned_action="ocr_overlay",
                            semantic_id=key,
                            occurrence_key=key,
                        )
                    )
                canvas.reading_order = [
                    obj.object_id for obj in objects if obj.canvas_id == canvas.canvas_id
                ]

        decision = resolve_profile(
            auto_suggestion=ContentProfile.PRESENTATION,
            confidence=0.9 if pictures or objects else 0.4,
            evidence=[f"slides:{len(prepared.canvases)}", f"pictures:{pictures}"],
            user_override=content_profile,
        )
        document = DocumentInfo(
            source_sha256=prepared.source_sha256,
            source_name=prepared.source_name,
            source_format=prepared.source_format,
            detected_mime=prepared.detected_mime,
            content_profile=decision.selected_profile,
            profile_source=decision.source,
            profile_confidence=decision.confidence,
            profile_evidence=decision.evidence,
            auto_profile_suggestion=decision.auto_suggestion,
            requested_mode=mode,
            selected_mode=selected,
            output_editability=(
                OutputEditability.RASTERIZED
                if selected is ProcessingMode.RENDERED
                else OutputEditability.EDITABLE
            ),
            input_asset=prepared.input_asset,
            derived_assets=derived_assets,
            conversion_lineage=list(prepared.conversion_lineage),
        )
        return DocumentStructureManifest(
            schema_version=CURRENT_SCHEMA_VERSION,
            manifest_id=build_manifest_id(prepared.source_sha256),
            created_at=datetime.now(timezone.utc),
            producer=ProducerInfo(
                name=PPTX_STRUCTURE_SCANNER_NAME,
                version=PPTX_STRUCTURE_SCANNER_VERSION,
            ),
            document=document,
            canvases=list(prepared.canvases),
            objects=objects,
            extensions={
                "picture_count": pictures,
                "slide_count": len(prepared.canvases),
                "physical_image_count": len(physical_image_keys),
                "physical_image_occurrence_count": pictures,
                "semantic_figure_count": 0,
                "table_count": sum(1 for item in objects if item.type is ObjectType.TABLE),
                "occurrence_count": sum(
                    1 for item in objects if item.type in {ObjectType.FIGURE, ObjectType.TABLE}
                ),
            },
        ).refresh_summary()
